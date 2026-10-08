# scripts/tflayout/27_paper_tables.py

import ast
import importlib.util
import itertools
import json
import math
import os
import re
import sys
import time

import numpy as np
import pandas as pd

CONFIG = dict(
    results_root="out/results",
    main_run="v8_headA_all",          # 论文主模型(第9批定)
    main_variant="perhead",           # 报告口径(各头各取最优权重)
    v3_run="v3_ce_marker_dense",      # 第9批之前的主线(S1 的参照、S2 位置机制的底)
    compare_main="out/results/_compare",          # 第9批 19 号：v8 5-seed headline + v8 对 v3 的 5 对 5
    head_a_main="out/results/_head_a_all",        # 第9批 25 号：v8 / v3 在 1108 个基因上
    ablation_compare=("out/results/_compare_b11", "out/results/_compare_b10"),    # 按顺序找；同一实验取 seed 最多的那份
    ablation_head_a=("out/results/_head_a_all_b11", "out/results/_head_a_all_b10"),
    head_a_seed_csv=("out/results/_head_a_all_b11/per_gene_head_a.csv", "out/results/_head_a_all_b10/per_gene_head_a.csv",
                     "out/results/_head_a_all/per_gene_head_a.csv"),   # 第一个有 v8 ≥4 个 seed 列的，给 Head A 全基因零分布
    baselines="out/results/_baselines",           # 26 号
    probe_dir="out/results/_siamese_probe",       # 2026-10-02a 第12批：28 号产出(没有就跳过)
    spec_dir="out/results/_spec_control",         # 2026-10-02b 第13批：29 号产出(没有就跳过)
    # 2026-10-02e 第16批：摘要里的常量(不是本脚本算的；来源 status 第3节 04 号 / 26 号 [0])。数据重建后要改这里
    fixed_facts=dict(n_tf=178, n_sites=214329, n_promoters=5373, source="status 第3节 04 号 / 26 号 [0]"),
    # (实验名, ASCII 标签, LaTeX 标签)；顺序就是表里的行顺序
    ablations=(("v10_abl_no_layout", "w/o layout", r"w/o layout ($L_g=\emptyset$)"),
               ("v10_abl_no_knockout", "w/o site deletion", r"w/o site deletion ($L_g\setminus D \to L_g$)"),
               ("v10_abl_no_cis", "w/o cis", r"w/o cis (promoter tokens)"),
               ("v10_lambda_a03", "lambda_A=0.3", r"$\lambda_A=0.3$")),
    pos_compare=dict(enabled=True, outdir="out/results/_compare_pos",
                     runs=("v3_ce_marker_dense", "v5_abl_no_position", "v5_abl_shuffle_position", "v6_abl_shift_position"),
                     labels=(("v5_abl_no_position", "w/o position", r"w/o position"),
                             ("v5_abl_shuffle_position", "shuffled positions", r"shuffled positions"),
                             ("v6_abl_shift_position", "global shift", r"global shift (spacing kept)")),
                     extra_pairs=(("v5_abl_no_position", "v6_abl_shift_position"),
                                  ("v5_abl_shuffle_position", "v6_abl_shift_position")),
                     n_boot=500),
    offset_grid=(-4.0, 4.0, 0.25),    # 跟 19 号一样(零分布里调 macro-F1 偏置用)
    outdir="out/results/_paper",
)


# ======================================================================================================
# 第16批(v6，2026-10-02e)新增：4~5 页短文(ICBCB/ICBBT)专用的紧凑表 + 自动摘要 + 堆叠规则。
# 都是模块级函数，输入是 run_paper_tables 里已经算好的 DataFrame，所以可以脱离真实数据单独测试(status 第5节第二十五版条目)。
# ======================================================================================================
_SHORT_COLS = (("B_r|TF", "B_r_within_tf", r"$r_{B|\mathrm{TF}}$", "r_B|TF"),
               ("C_AUCdn", "C_auroc_down", r"AUROC$_\downarrow$", "AUROC_dn"),
               ("C_APdn", "C_auprc_down", r"AUPRC$_\downarrow$", "AUPRC_dn"),
               ("C_APup", "C_auprc_up", r"AUPRC$_\uparrow$", "AUPRC_up"),
               ("wTF_dn", "wtf_dn", r"wTF-AUROC$_\downarrow$", "wTF-AUROC_dn"))
_SHORT_ROWS = (("B4_pair+TF先验", "B4 pair + TF prior"),
               ("B6_gbdt(同B5特征)", "B6 GBDT + WT expr."),
               ("B7a_gbdt+flat(无WT)", "B7a GBDT + flat layout"),
               ("B7a_bag(无WT)", "B7a-bag (bagged)"),
               ("B7b_bag(+WT)", "B7b-bag (+ WT expr.)"))
_MECH_COLS = (("B_r_within_tf", r"$\Delta r_{B|\mathrm{TF}}$", "d r_B|TF"),
              ("C_auprc_down", r"$\Delta$AUPRC$_\downarrow$", "d AUPRC_dn"),
              ("C_auprc_up", r"$\Delta$AUPRC$_\uparrow$", "d AUPRC_up"))
_METRIC_PHRASE = {"B_r|TF": "within-TF log2FC correlation", "C_APdn": "down-regulation AUPRC",
                  "wTF_dn": "within-TF down-regulation AUROC"}


def _fnum(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return float("nan")


def _fv(v, sign=False, nd=3):
    v = _fnum(v)
    if not np.isfinite(v):
        return "--"
    return f"{v:+.{nd}f}" if sign else f"{v:.{nd}f}"


def _istrue(x):
    return isinstance(x, (bool, np.bool_)) and bool(x)


def _tex_cell(c):
    """c = (文本, 标记字符串(†‡§ 的任意组合), 是否粗体)。数字(带符号或 ±)放进数学模式。"""
    txt, marks, bold = c
    if txt == "--":
        return "--"
    mk = marks.replace("†", r"\dagger").replace("‡", r"\ddagger").replace("§", r"\S")
    if txt[:1] in "+-" or "±" in txt:
        s = "$" + txt.replace("±", r"\pm ") + (f"^{{{mk}}}$" if mk else "$")
    else:
        s = txt + (f"$^{{{mk}}}$" if mk else "")
    return rf"\textbf{{{s}}}" if bold else s


def _plain_cell(c):
    txt, marks, bold = c
    return txt + marks


def _cell(v, marks="", bold=False, sign=False, nd=3):
    return (_fv(v, sign, nd), marks if np.isfinite(_fnum(v)) else "", bold)


def render_short_table(tbl, tex_table, sink):
    """tbl = dict(caption, label, header_tex, header_txt, rows, notes)；rows 的元素：('rule',) 或 ('row', 标签tex, 标签txt, [cell...])。
    返回纯文本行(给 summary.txt / .md 用)，同时把 LaTeX 追加进 sink。"""
    trows, plain = [], []
    for r_ in tbl["rows"]:
        if r_[0] == "rule":
            trows.append(r"\midrule")
            plain.append("-" * 60)
            continue
        _, lab_tex, lab_txt, cells = r_
        trows.append(lab_tex + " & " + " & ".join(_tex_cell(c) for c in cells))
        plain.append(lab_txt.ljust(44) + "".join(f"{_plain_cell(c):>13}" for c in cells))
    tex_table(tbl["caption"], tbl["label"], tbl["header_tex"], trows, "l" + "c" * (len(tbl["header_tex"]) - 1),
              notes=tbl.get("notes", ""), wide=False, sink=sink)
    return ["".ljust(44) + "".join(f"{h:>13}" for h in tbl["header_txt"][1:])] + plain


def write_short_docx(path, tables):
    """把 tables(同 render_short_table 的输入)写成一个 Word 文件，三线表，Times New Roman 8pt，方便粘进会议的 Word 模板。
    没有 python-docx 就返回 False(其余输出照常)。"""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError:
        return False

    def border(cell, **kw):
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders")
            tcPr.append(b)
        for edge, sz in kw.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
            b.append(el)

    def put(cell, text, bold=False, left=False, italic=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(text)
        run.font.size = Pt(8)
        run.font.name = "Times New Roman"
        run.bold = bold
        run.italic = italic

    doc = Document()
    for t_ in tables:
        cp = doc.add_paragraph()
        r0 = cp.add_run(t_["caption_txt"])
        r0.font.size = Pt(9)
        r0.font.name = "Times New Roman"
        header = t_["header_txt"]
        rows = [r_ for r_ in t_["rows"] if r_[0] != "rule"]
        tb_ = doc.add_table(rows=1 + len(rows), cols=len(header))
        tb_.autofit = False
        w0 = 2.0
        wr = (6.0 - w0) / max(len(header) - 1, 1)
        for ci_, col_ in enumerate(tb_.columns):
            for cell_ in col_.cells:
                cell_.width = Inches(w0 if ci_ == 0 else wr)
        for j, h in enumerate(header):
            put(tb_.rows[0].cells[j], h, bold=True, left=(j == 0))
            border(tb_.rows[0].cells[j], top=12, bottom=6)
        i_ = 0
        for r_ in t_["rows"]:
            if r_[0] == "rule":
                if i_ > 0:   # 分组线：给上一行加下边线
                    for c_ in tb_.rows[i_].cells:
                        border(c_, bottom=4)
                continue
            _, _, lab_txt, cells = r_
            i_ += 1
            put(tb_.rows[i_].cells[0], lab_txt, left=True)
            for j, c in enumerate(cells, 1):
                put(tb_.rows[i_].cells[j], _plain_cell(c), bold=c[2])
        for c_ in tb_.rows[-1].cells:
            border(c_, bottom=12)
        if t_.get("notes_txt"):
            npar = doc.add_paragraph()
            rn = npar.add_run(t_["notes_txt"])
            rn.font.size = Pt(7.5)
            rn.font.name = "Times New Roman"
        doc.add_paragraph()
    doc.save(path)
    return True


def short_main_table(T2, mdl26, T6, t6v, facts_test, n_seeds):
    """表一：主比较。行 = 代表性基线 + 我们的单 seed(均值±sd)/集成/堆叠 + (i)(j) 两行差值。返回 tbl 或 None。"""
    if T2 is None or T2.empty or not mdl26:
        return None
    by_model = {r_["model"]: r_ for r_ in T2.to_dict("records")}
    ours = by_model.get(mdl26)
    if ours is None:
        return None
    comp = [by_model[k] for k, _ in _SHORT_ROWS if k in by_model] + [ours]
    best = {c: np.nanmax([_fnum(r_.get(c, np.nan)) for r_ in comp]) for _, c, _, _ in _SHORT_COLS}
    rows = []
    for key, lab in _SHORT_ROWS:
        r_ = by_model.get(key)
        if r_ is None:
            continue
        rows.append(("row", lab, lab, [_cell(r_.get(c), bold=abs(_fnum(r_.get(c)) - best[c]) < 5e-4)
                                       for _, c, _, _ in _SHORT_COLS]))
    rows.append(("rule",))
    # 我们的单 seed 均值 ± sd(26 号 table6_values.csv)
    sing = None
    if t6v is not None and len(t6v):
        qs = t6v[t6v["label"].astype(str).str.startswith("模型") & t6v["label"].astype(str).str.contains("单seed")]
        if len(qs):
            sing = qs.set_index("metric")
    if sing is not None:
        cells = []
        for tag, _, _, _ in _SHORT_COLS:
            if tag in sing.index:
                v_, sd_ = _fnum(sing.loc[tag, "value"]), _fnum(sing.loc[tag].get("seed_sd", np.nan))
                cells.append((f"{v_:.3f}" + (f"±{sd_:.3f}" if np.isfinite(sd_) else ""), "", False))
            else:
                cells.append(("--", "", False))
        rows.append(("row", r"Ours, 1 seed (mean $\pm$ sd)", "Ours, 1 seed (mean ± sd)", cells))
    rows.append(("row", rf"\textbf{{Ours, {n_seeds}-seed ens.}}", f"Ours, {n_seeds}-seed ens.",
                 [_cell(ours.get(c), bold=abs(_fnum(ours.get(c)) - best[c]) < 5e-4) for _, c, _, _ in _SHORT_COLS]))
    stk = by_model.get("STACK(模型+B7b_bag)")
    if stk is not None:
        rows.append(("row", r"\textit{Ours $+$ B7b-bag (stack)}", "Ours + B7b-bag (stack)",
                     [_cell(stk.get(c)) for _, c, _, _ in _SHORT_COLS]))
    # (i)(j) 差值行
    if T6 is not None and not T6.empty:
        has_tot = "tot_sig" in T6.columns
        rows.append(("rule",))
        for key, lab_tex, lab_txt in (("(i) 集成−B7a_bag(对等集成)", r"Ours ens.\ $-$ B7a-bag", "Ours ens. - B7a-bag"),
                                      ("(j) 单seed均值−B7a(单对单)", r"Ours 1 seed $-$ B7a", "Ours 1 seed (mean) - B7a")):
            q = T6[T6["comparison"] == key].set_index("metric")
            if q.empty:
                continue
            cells = []
            for tag, _, _, _ in _SHORT_COLS:
                if tag not in q.index:
                    cells.append(("--", "", False))
                    continue
                r_ = q.loc[tag]
                dag = bool(r_["ci_lo"] > 0 or r_["ci_hi"] < 0)
                ddg = bool(has_tot and _istrue(r_.get("tot_sig", False)))
                cells.append(_cell(r_["delta"], ("†" if dag else "") + ("‡" if ddg else ""), sign=True))
            rows.append(("row", lab_tex, lab_txt, cells))
    nrow, ngene, nsig = facts_test["rows"], facts_test["genes"], facts_test["sig"]
    cap = (rf"Held-out-chromosome test ({nrow:,} gene$\times$TF pairs, {ngene} genes, {nsig:,} significant). "
           rf"Top: absolute values; bold: best among the baselines and our ensemble. Bottom: differences; "
           rf"$\dagger$: 95\% gene-cluster bootstrap CI excludes 0; $\ddagger$: the interval that also includes re-training variability excludes 0.")
    notes = (r"B7: gradient boosting on the same layout flattened into features; bag: bagged GBDT members, each on a random 80\% of the "
             r"training genes. B6/B7b use measured wild-type expression (more input information than our model). "
             r"wTF-AUROC: pooled within-TF AUROC.")
    return dict(caption=cap, label="tab:short_main", header_tex=["Method"] + [h for _, _, h, _ in _SHORT_COLS],
                header_txt=["Method"] + [h for _, _, _, h in _SHORT_COLS], rows=rows, notes=notes,
                caption_txt=(f"Table 1. Main comparison on held-out chromosomes ({nrow:,} gene-TF pairs, {ngene} genes, {nsig:,} significant). "
                             "Top: absolute values (bold: best among baselines and our ensemble). Bottom: differences; "
                             "† bootstrap 95% CI excludes 0; ‡ the interval that also includes re-training variability excludes 0."),
                notes_txt=("B7: gradient boosting on the same layout flattened into features; bag: bagged GBDT members, each on a random 80% "
                           "of the training genes. B6/B7b use measured wild-type expression. wTF-AUROC: pooled within-TF AUROC."))


def short_mech_table(T4, T5, ablations, spec_df, d_in_lg=None):
    """表二：机制。按 (k, 参照取值) 分组：每组一行参照(绝对值)，组内放该组的消融；最后一行是推理时删错 TF。列 = ΔB_r|TF、ΔAUPRC dn/up。
    (v6.1：v6 只放了一行参照，取的是第一个消融(k=2)的，对 k=5 的 w/o site deletion 是错的。)"""
    if (T4 is None or T4.empty) and spec_df is None:
        return None

    def mk_of(r_, k_, with_cons=True):
        return ("†" if _istrue(r_["test_sig"]) else "") + ("‡" if _istrue(r_.get("tot_sig", False)) else "") + \
               ("§" if with_cons and k_ > 2 and _istrue(r_.get("cons_sig", False)) else "")

    def ref_tuple(run):
        out = []
        for met, _, _ in _MECH_COLS:
            q = T4[(T4["run"] == run) & (T4["metric"] == met) & (T4["scope"] == "test")]
            v_ = _fnum(q["ref"].iloc[0]) if len(q) else float("nan")
            out.append(round(v_, 6) if np.isfinite(v_) else None)   # 缺失用 None(NaN 不能当字典键)
        return tuple(out)

    entries = []   # (key, [row, ...])；key = (k, 参照取值)
    if T4 is not None and not T4.empty:
        for run, tex_lab, txt_lab in (("v10_abl_no_layout", r"w/o layout ($L_g=\emptyset$)", "w/o layout"),
                                      ("v10_abl_no_knockout", r"w/o site deletion", "w/o site deletion"),
                                      ("v10_abl_no_cis", r"w/o cis", "w/o cis")):
            if T4[T4["run"] == run].empty:
                continue
            cells, k_ = [], None
            for met, _, _ in _MECH_COLS:
                q = T4[(T4["run"] == run) & (T4["metric"] == met) & (T4["scope"] == "test")]
                if not len(q) or not np.isfinite(_fnum(q["delta"].iloc[0])):
                    cells.append(("--", "", False))
                    continue
                r_ = q.iloc[0]
                k_ = int(r_["k"])
                cells.append(_cell(r_["delta"], mk_of(r_, k_, with_cons=(run != "v10_abl_no_cis")), sign=True))
            if k_ is None:
                continue
            rws = [("row", tex_lab + rf" ($k={k_}$)", txt_lab + f" (k={k_})", cells)]
            if run == "v10_abl_no_knockout" and T5 is not None and not T5.empty:
                cells2, k2 = [], k_
                for met, _, _ in _MECH_COLS:
                    q = T5[(T5["kind"] == "消融−参照") & (T5["against"] == "w/o site deletion") &
                           (T5["stratum"] == "D∈L_g") & (T5["metric"] == met)]
                    if not len(q):
                        cells2.append(("--", "", False))
                        continue
                    r_ = q.iloc[0]
                    k2 = int(r_["k"])
                    cells2.append(_cell(r_["delta"], mk_of(r_, k2), sign=True))
                rws.append(("row", r"\quad only $D\in L_g$ pairs" + rf" ($k={k2}$)", "   only D in L_g pairs" + f" (k={k2})", cells2))
            entries.append(((k_, ref_tuple(run)), rws))
    rows, seen = [], []
    for key, _ in entries:
        if key not in seen:
            seen.append(key)
    for gi, key in enumerate(seen):
        if gi:
            rows.append(("rule",))
        rows.append(("row", rf"Reference ($k={key[0]}$ seeds)", f"Reference (k={key[0]} seeds)", [_cell(v_) for v_ in key[1]]))
        for key2, rws in entries:
            if key2 == key:
                rows += rws
    if spec_df is not None and len(spec_df):
        q0 = spec_df[spec_df["contrast"] == "Δ_spec"].set_index("metric")
        cells = []
        for met, _, _ in _MECH_COLS:
            if met in q0.index:
                r_ = q0.loc[met]
                cells.append(_cell(r_["delta"], "†" if _istrue(r_["test_sig"]) else "", sign=True))
            else:
                cells.append(("--", "", False))
        if rows:
            rows.append(("rule",))
        rows.append(("row", r"Wrong $-$ correct TF deleted", "Wrong - correct TF deleted (inference)", cells))
    if not [r_ for r_ in rows if r_[0] == "row"]:
        return None
    cap = (r"Mechanism: component ablations ($\Delta$ = ablation $-$ reference) and a deletion-specificity control at inference. "
           r"Reference rows: the full model trained with the same seeds, absolute values on all test pairs ($k$ = number of seeds). "
           r"$\dagger$: 95\% gene-cluster bootstrap CI excludes 0; $\ddagger$: the interval that also includes re-training variability excludes 0; "
           r"$\S$ ($k>2$): also excludes 0 without extrapolating the 2-seed variability to $k$ seeds (conservative).")
    notes = (r"Last row (inference, $D\in L_g$ pairs): delete the sites of another TF present in the promoter instead of $D$'s; same trained weights for both inputs "
             r"(no re-training variability, hence $\dagger$ only). $D\in L_g$: the depleted TF $D$ has a mapped site in the gene's promoter.")
    out_ = dict(caption=cap, label="tab:short_mech", header_tex=["Variant"] + [h for _, h, _ in _MECH_COLS],
                header_txt=["Variant"] + [h for _, _, h in _MECH_COLS], rows=rows, notes=notes,
                caption_txt=("Table 2. Mechanism: component ablations (Δ = ablation − reference) and a deletion-specificity control at inference. "
                             "Reference rows: the full model trained with the same seeds, absolute values on all test pairs (k = number of seeds). "
                             "† bootstrap 95% CI excludes 0; ‡ the interval that also includes re-training variability excludes 0; "
                             "§ (k>2) also excludes 0 without extrapolating the 2-seed variability to k seeds."),
                notes_txt=("Last row (inference, D in L_g pairs): delete the sites of another TF present in the promoter instead of D's; same trained weights "
                           "for both inputs (no re-training variability, hence † only). D in L_g: the depleted TF D has a mapped site in the gene's promoter."))
    if d_in_lg:   # v6.2：子集行的 Δ 是在子集内算的
        out_["notes"] += rf" The row ``only $D\in L_g$ pairs'' is computed within the {100 * d_in_lg:.1f}\% of test pairs in which $D$ has a site."
        out_["notes_txt"] += f" The 'only D in L_g pairs' row is computed within the {100 * d_in_lg:.1f}% of test pairs in which D has a site."
    return out_


def build_abstract(T2, mdl26, T4, T5, T6, t5b, facts, spec_df, fixed, n_seeds):
    """按 (m)(第9节(l)3)选摘要版本并自动填数。返回 dict(text, variant, flags, sources, words)。flags 非空 = 有一句的条件没满足，写论文前先看。"""
    flags, srcs = [], []
    ft = facts["test"]
    by_model = {r_["model"]: r_ for r_ in T2.to_dict("records")} if T2 is not None and not T2.empty else {}
    ours = by_model.get(mdl26, {})
    a_dn, a_up = _fnum(ours.get("C_auroc_down")), _fnum(ours.get("C_auroc_up"))
    p_dn, p_up = _fnum(ours.get("C_auprc_down")), _fnum(ours.get("C_auprc_up"))
    b_r = _fnum(ours.get("B_r"))
    srcs += [(f"{a_dn:.3f}/{a_up:.3f}", "Table 2 模型行 C_AUCdn/C_AUCup(table2_head_bc.csv)"),
             (f"{p_dn:.3f}/{p_up:.3f}", "Table 2 模型行 C_APdn/C_APup"), (f"{b_r:.3f}", "Table 2 模型行 B_r"),
             (f"{ft['rows']:,}/{ft['tfs']}", "[0] test 对数/被耗竭 TF 数(主模型 17 号导出)"),
             (f"{fixed['n_tf']}/{fixed['n_sites']:,}/{fixed['n_promoters']:,}", "常量：" + fixed.get("source", "status 第3节 04 号"))]
    if not all(np.isfinite([a_dn, a_up, p_dn, p_up, b_r])):
        flags.append("模型行的 AUROC/AUPRC/B_r 有缺失(Table 2 没生成)，摘要第三句的数字不可信")
    # ---- 比较句：(m)
    variant, items, items_c, jitems = "UNKNOWN", [], [], []
    if T6 is not None and not T6.empty and "tot_lo" in T6.columns:
        def g6t(key, met):
            q_ = T6[(T6["comparison"] == key) & (T6["metric"] == met)]
            if not len(q_):
                return (np.nan,) * 5
            r0 = q_.iloc[0]
            return tuple(_fnum(r0[c_]) for c_ in ("delta", "tot_lo", "tot_hi", "cons_lo", "cons_hi"))
        got_i = {m_: g6t("(i) 集成−B7a_bag(对等集成)", m_) for m_ in ("B_r|TF", "C_APdn", "wTF_dn")}
        got_j = {m_: g6t("(j) 单seed均值−B7a(单对单)", m_) for m_ in ("B_r|TF", "C_APdn", "wTF_dn")}
        items = [m_ for m_, v_ in got_i.items() if np.isfinite(v_[1]) and v_[0] >= 0.02 and v_[1] > 0]
        items_c = [m_ for m_ in items if np.isfinite(got_i[m_][3]) and got_i[m_][3] > 0]
        jitems = [m_ for m_, v_ in got_j.items() if np.isfinite(v_[1]) and v_[0] >= 0.02 and v_[1] > 0]
        variant = "A" if len(items) >= 2 else ("A-lite" if len(items) == 1 else "B")
        if variant in ("A", "A-lite"):
            parts = [f"{_METRIC_PHRASE[m_]} (+{got_i[m_][0]:.3f})" for m_ in items]
            for m_ in items:
                srcs.append((f"+{got_i[m_][0]:.3f}", f"Table 6 (i) 集成−B7a_bag {m_}，合成区间 ({got_i[m_][1]:+.3f},{got_i[m_][2]:+.3f})"
                             + ("；§ 也成立" if m_ in items_c else "；§ 不成立")))
            if jitems:
                flags.append(f"(j‡) 有明显项 {jitems}：\"matches single models\" 不成立，比较句要重写")
            if variant == "A-lite":
                cmp_s = ("Gradient boosting on the same layout flattened into features matches single models; under matched ensembling "
                         f"the siamese model retains a modest advantage in {parts[0]}.")
            else:
                cmp_s = ("Gradient boosting on the same layout flattened into features matches single models, but under matched ensembling "
                         f"the siamese model improves {' and '.join(parts)} (intervals that include re-training variability exclude 0).")
            if any(m_ not in items_c for m_ in items):
                flags.append(f"比较句里的 {[m_ for m_ in items if m_ not in items_c]} 过 ‡ 但不过 §(外推不稳)，措辞用 \"modest\" 并在正文说明")
        else:
            cmp_s = ("Gradient boosting on the same layout flattened into features performs comparably, both as single models "
                     "and as matched ensembles.")
    else:
        flags.append("Table 6 没有 ‡ 列(26 号不是 v6 的产出)：(m) 无法判定，比较句暂用 Variant B 的保守写法")
        cmp_s = ("Gradient boosting on the same layout flattened into features performs comparably, both as single models "
                 "and as matched ensembles.")
    # ---- 消融/删位点
    lay_ok = False
    if T4 is not None and not T4.empty:
        q = T4[(T4["run"] == "v10_abl_no_layout") & (T4["scope"] == "test") & T4["metric"].isin(["B_r", "B_r_within_tf", "C_auprc_down"])]
        lay_ok = bool(len(q) == 3 and (q["delta"] < 0).all() and q["tot_sig"].map(_istrue).all())
    if not lay_ok:
        flags.append("no_layout 在 B_r/B_r|TF/C_APdn 上不是三项都 ‡ 变差：\"layout 是主要信息来源\" 这句要改")
    dn_s = up_s = None
    if T5 is not None and not T5.empty:
        def d5(met):
            q = T5[(T5["kind"] == "消融−参照") & (T5["against"] == "w/o site deletion") & (T5["stratum"] == "D∈L_g") & (T5["metric"] == met)]
            return q.iloc[0] if len(q) else None
        rdn, rup = d5("C_auprc_down"), d5("C_auprc_up")
        if rdn is not None and rup is not None:
            sd_ok = all(_istrue(r_["tot_sig"]) and _istrue(r_.get("cons_sig", False)) and r_["delta"] < 0 for r_ in (rdn, rup))
            dn_s, up_s = -_fnum(rdn["delta"]), -_fnum(rup["delta"])
            if not (dn_s > 0 and up_s > 0):
                flags.append("删位点在 D∈L_g 层不是两项都变差(消融没有更差)：删位点那句整句去掉")
                dn_s = up_s = None
        if dn_s is not None:
            srcs += [(f"{dn_s:.3f}/{up_s:.3f}", "Table 5 w/o site deletion, D∈L_g 层 C_APdn/C_APup(消融−参照取负)"
                      + ("；‡§ 都成立" if sd_ok else "；⚠ 不是都过 ‡§"))]
            if not sd_ok:
                flags.append("删位点在 D∈L_g 层的 C_APdn/C_APup 不是都过 ‡§：\"improves ... by\" 要降成 \"in a test-sampling sense\" 或去掉数字")
    if dn_s is None:
        flags.append("没有 Table 5 的 w/o site deletion D∈L_g 层：删位点那句的数字缺")
    sp_dn = sp_up = None
    if spec_df is not None and len(spec_df):
        q0 = spec_df[spec_df["contrast"] == "Δ_spec"].set_index("metric")
        if {"C_auprc_down", "C_auprc_up"} <= set(q0.index):
            sp_dn, sp_up = -_fnum(q0.loc["C_auprc_down", "delta"]), -_fnum(q0.loc["C_auprc_up", "delta"])
            ok_sp = all(_istrue(q0.loc[m_, "test_sig"]) and q0.loc[m_, "delta"] < 0 for m_ in ("C_auprc_down", "C_auprc_up"))
            if not (sp_dn > 0 and sp_up > 0):
                sp_dn = sp_up = None
            else:
                srcs.append((f"{sp_dn:.3f}/{sp_up:.3f}", "29 号 contrasts_spec.csv Δ_spec C_APdn/C_APup(取负)" + ("；† 成立" if ok_sp else "；⚠ 不成立")))
            if not ok_sp:
                flags.append("29 号 (q1) 不成立：\"deleting another TF's sites lowers ...\" 这半句要删")
    if sp_dn is None:
        flags.append("29 号 Δ_spec 缺失(先 --spec-control 再 --tables)或不是两项都变差：\"deleting another TF's sites\" 那半句不写")
    # ---- 互补性、WT 表达
    wt_ok = False
    if T6 is not None and not T6.empty:
        q = T6[(T6["comparison"] == "集成−B7b_bag") & (T6["metric"] == "B_r|TF")]
        wt_ok = bool(len(q) and q["delta"].iloc[0] < 0 and (q["ci_hi"].iloc[0] < 0))
    if not wt_ok:
        flags.append("模型集成−B7b_bag 的 B_r|TF 不是 † 变差：\"Measured wild-type expression adds more than model structure\" 要改")
    st_dn = st_up = None
    st_ok = False
    if t5b is not None and len(t5b):
        q = t5b[(t5b["stack"].astype(str) == "STACK(模型+B7b_bag)") & (t5b["reference"].astype(str) == "B7b_bag(+WT)")].set_index("metric")
        if {"C_APdn", "C_APup"} <= set(q.index):
            st_dn, st_up = _fnum(q.loc["C_APdn", "delta"]), _fnum(q.loc["C_APup", "delta"])
            st_ok = bool(q.loc["C_APdn", "ci_lo"] > 0 and q.loc["C_APup", "ci_lo"] > 0)
            srcs.append((f"{st_dn:.3f}/{st_up:.3f}", "26 号 table5_stack.csv 堆叠(模型+B7b_bag)−B7b_bag 的 C_APdn/C_APup"
                         + ("；† 都成立" if st_ok else "；⚠ 不是都 †")))
    if st_dn is None:
        flags.append("没有 STACK(模型+B7b_bag) 对 B7b_bag 的 C_APdn/C_APup：互补性那句的数字缺")
    elif not st_ok:
        flags.append("堆叠对 B7b_bag 的 C_APdn/C_APup 不是都 †：互补性那句要降级")
    # ---- 拼摘要
    s = [("Predicting how genes respond when a transcription factor (TF) is depleted is hard because TF binding and regulation "
          "overlap only partially."),
         (f"We ask how much a binding layout of {fixed['n_tf']} TFs ({fixed['n_sites']:,} mapped sites in "
          f"{fixed['n_promoters']:,} promoters, from ChEC-seq) explains depletion responses on held-out chromosomes "
          f"({ft['rows']:,} gene-TF pairs, {ft['tfs']} depletions)."),
         (f"A siamese multi-task model that contrasts the wild-type layout with the layout after removing the depleted TF's sites reaches "
          f"direction AUROC {a_dn:.3f}/{a_up:.3f} (down/up), AUPRC {p_dn:.3f}/{p_up:.3f} and log2FC r = {b_r:.3f} as a {n_seeds}-seed ensemble."),
         cmp_s]
    clause = []
    if lay_ok:
        clause.append("Ablations identify the layout as the dominant information source")
    if dn_s is not None:
        clause.append(f"training with site deletion improves down- and up-regulation AUPRC by {dn_s:.3f} and {up_s:.3f} on genes where "
                      "the depleted TF has a site")
    if clause:
        t_ = "; ".join(clause)
        t_ = t_[0].upper() + t_[1:]
        if sp_dn is not None:
            t_ += f", and deleting another TF's sites instead lowers AUPRC by {sp_dn:.3f} and {sp_up:.3f}"
        s.append(t_ + ".")
    tail = []
    if wt_ok:
        tail.append("Measured wild-type expression adds more than model structure")
    if st_dn is not None and st_ok:
        tail.append(f"the model remains complementary to the strongest baseline (stacking raises AUPRC by {st_dn:.3f}/{st_up:.3f})")
    if tail:
        s.append((", yet ".join(tail) if len(tail) == 2 else tail[0]) + ".")
    s.append("We also document failure modes: unseen TFs and spacing below ~100 bp.")
    text = " ".join(s)
    return dict(text=text, variant=variant, items=items, items_c=items_c, flags=flags, sources=srcs, words=len(text.split()))


def rules_stack(t5b, clear):
    """(f1)/(f2) 按堆叠分开判：每个堆叠 × 每个参照一行，带标签(旧版把几个堆叠的数混进一行，且 (f1) 只看第一个基线)。"""
    out = []
    if t5b is None or not len(t5b):
        return out
    for stk in dict.fromkeys(t5b["stack"].astype(str)):
        qs = t5b[(t5b["stack"].astype(str) == stk) & t5b["metric"].isin(["B_r|TF", "C_APdn"])]
        refs = list(dict.fromkeys(qs["reference"].astype(str)))
        for tag, desc, rr in (("(f1)", "堆叠−基线 明显>0(模型带来基线之外的信息)", [x for x in refs if not x.startswith("模型")]),
                              ("(f2)", "堆叠−模型 明显>0(基线特征里有模型缺的信息)", [x for x in refs if x.startswith("模型")])):
            for ref in rr:
                q = qs[qs["reference"].astype(str) == ref]
                out.append((f"{tag} {stk} − {ref}：{desc}",
                            any(clear(r_.delta, r_.ci_lo, r_.ci_hi) and r_.delta > 0 for r_ in q.itertuples()),
                            "  ".join(f"{r_.metric} {r_.delta:+.3f}[{r_.ci_lo:+.3f},{r_.ci_hi:+.3f}]" for r_ in q.itertuples())))
    return out


def run_paper_tables(results_root, main_run, main_variant, v3_run, compare_main, head_a_main, ablation_compare,
                     ablation_head_a, head_a_seed_csv, baselines, ablations, pos_compare, offset_grid, outdir,
                     probe_dir=None, spec_dir=None, fixed_facts=None):
    """唯一入口，各段见文件头。"""
    t_all = time.time()
    here = os.path.dirname(os.path.abspath(__file__))
    os.makedirs(os.path.join(outdir, "fig"), exist_ok=True)
    lines, md, tex = [], [], []
    cls = {"down": 0, "ns": 1, "up": 2}
    M_LONG = ("A_r_gene", "B_r", "B_sign_acc", "B_r_within_tf", "C_auroc_down", "C_auroc_up", "C_auprc_down",
              "C_auprc_up", "C_macro_f1_offset")
    SHORT = dict(A_r_gene="A_r", B_r="B_r", B_sign_acc="B_sign", B_r_within_tf="B_r|TF", C_auroc_down="C_AUCdn",
                 C_auroc_up="C_AUCup", C_auprc_down="C_APdn", C_auprc_up="C_APup", C_macro_f1_offset="C_mF1",
                 B_dr="B_dr", B_dr_ns="B_drns", A_all="A_r(1108)")
    TEXM = dict(A_r_gene=r"$r_A$ (grid)", A_all=r"$r_A$ (all)", B_r=r"$r_B$", B_sign_acc="Sign acc.",
                B_r_within_tf=r"$r_B$ within TF", C_auroc_down=r"AUROC$_\downarrow$", C_auroc_up=r"AUROC$_\uparrow$",
                C_auprc_down=r"AUPRC$_\downarrow$", C_auprc_up=r"AUPRC$_\uparrow$", C_macro_f1_offset="Macro-F1",
                wtf_dn=r"wTF-AUROC$_\downarrow$", wtf_up=r"wTF-AUROC$_\uparrow$")
    B26 = {"B_r": "B_r", "B_sign": "B_sign_acc", "B_r|TF": "B_r_within_tf", "C_AUCdn": "C_auroc_down",
           "C_AUCup": "C_auroc_up", "C_APdn": "C_auprc_down", "C_APup": "C_auprc_up", "C_mF1": "C_macro_f1_offset"}
    NAME = {"B1_TF均值": ("B1 TF mean", "B1: TF mean"),
            "B2_TF×结合先验": ("B2 TF x binding prior", r"B2: TF $\times$ binding prior"),
            "B3_gene_only": ("B3 gene-only", "B3: gene features only"),
            "B4_pair+TF先验": ("B4 pair+TF prior", "B4: pair features + TF prior"),
            "B5_B4+实测WT表达": ("B5 B4+WT expr", "B5: B4 + measured WT expr."),
            "B6_gbdt(同B5特征)": ("B6 GBDT(B5 feats)", "B6: GBDT (B5 features)"),
            "B7a_gbdt+flat(无WT)": ("B7a GBDT+flat layout", "B7a: GBDT, B4 feats + flat layout (no WT expr.)"),
            "B7a_bag(无WT)": ("B7a-bag GBDT+flat (bagged)", "B7a-bag: bagged GBDTs, B7a features"),
            "B7b_gbdt+flat(+WT)": ("B7b GBDT+flat layout+WT", "B7b: GBDT, B7a + measured WT expr."),
            "B7b_bag(+WT)": ("B7b-bag GBDT+flat+WT (bagged)", "B7b-bag: bagged GBDTs, B7b features"),
            "ORACLE_同基因其它TF(参照)": ("Oracle same gene", "Oracle: same gene, other TFs"),
            "A1_layout_ridge": ("A1 layout ridge", "A1: layout counts, ridge"),
            "A2_token_ridge": ("A2 token ridge", "A2: BPE-token bag, ridge"),
            "A3_layout+token_分块ridge": ("A3 layout+token", "A3: A1+A2, block ridge"),
            "A2k_6mer_ridge": ("A2k 6-mer ridge", "A2k: 6-mer counts, ridge"),
            "A3k_layout+6mer_分块ridge": ("A3k layout+6-mer", "A3k: layout + 6-mer, block ridge"),
            "A4_gbdt(layout+token)": ("A4 GBDT layout+token", "A4: GBDT (layout + tokens)"),
            "A4k_gbdt(layout+6mer)": ("A4k GBDT layout+6-mer", "A4k: GBDT (layout + 6-mer)"),
            "CITRA": ("CITRA", "CITRA")}

    def say(msg=""):
        print(msg, flush=True)
        lines.append(str(msg))

    def pearson(x, y):
        x, y = np.asarray(x, np.float64), np.asarray(y, np.float64)
        ok = np.isfinite(x) & np.isfinite(y)
        x, y = x[ok], y[ok]
        if len(x) < 3 or x.std() <= 1e-12 * max(1.0, abs(float(x.mean()))) or y.std() == 0:
            return float("nan")  # 近常数也当常数(跟 25 号第10批同一规则)
        return float(np.corrcoef(x, y)[0, 1])

    def auroc_auprc(score, pos):  # 跟 19 号逐行相同
        pos = np.asarray(pos, bool)
        n1 = int(pos.sum())
        n0 = len(pos) - n1
        if n1 == 0 or n0 == 0:
            return float("nan"), float("nan")
        sc = np.asarray(score, np.float64)
        _, inv_, cnt_ = np.unique(sc, return_inverse=True, return_counts=True)
        avg_rank = np.cumsum(cnt_) - (cnt_ - 1) / 2.0
        r = avg_rank[inv_.reshape(-1)]
        auc = float((r[pos].sum() - n1 * (n1 + 1) / 2.0) / (n1 * n0))
        p = pos[np.argsort(-sc, kind="mergesort")]
        prec = np.cumsum(p) / np.arange(1, len(p) + 1)
        return auc, float(prec[p].sum() / n1)

    def macro_f1(pred, true):
        cm = np.bincount(true * 3 + pred, minlength=9).reshape(3, 3)
        f1 = []
        for c in range(3):
            tp, fp, fn = cm[c, c], cm[:, c].sum() - cm[c, c], cm[c, :].sum() - cm[c, c]
            f1.append(2 * tp / (2 * tp + fp + fn) if (2 * tp + fp + fn) else 0.0)
        return float(np.mean(f1))

    grid_off = np.arange(offset_grid[0], offset_grid[1] + 1e-9, offset_grid[2])

    def tune(prob_val, yv_):
        lp = np.log(np.clip(prob_val, 1e-12, 1.0))
        best = (-1.0, 0.0, 0.0)
        for od in grid_off:
            for ou in grid_off:
                f = macro_f1((lp + np.array([od, 0.0, ou])).argmax(1), yv_)
                if f > best[0] + 1e-12:
                    best = (f, float(od), float(ou))
        return best

    def apply_off(prob, best):
        return (np.log(np.clip(prob, 1e-12, 1.0)) + np.array([best[1], 0.0, best[2]])).argmax(1)

    def rd(path):
        if path and os.path.exists(path):
            try:
                return pd.read_csv(path)
            except Exception as e:  # noqa: BLE001 —— 读不出来就当缺失，段落里会提示
                say(f"  ⚠ 读 {path} 失败：{type(e).__name__}: {e}")
        return None

    def rjson(path):
        if path and os.path.exists(path):
            with open(path, encoding="utf-8") as fh:
                return json.load(fh)
        return None

    def mtime(path):
        return time.strftime("%Y-%m-%d %H:%M", time.localtime(os.path.getmtime(path))) if os.path.exists(path) else "缺"

    def f3(v, sign=False):
        if v is None or not np.isfinite(v):
            return "nan"
        return f"{v:+.3f}" if sign else f"{v:.3f}"

    def texesc(s):
        return str(s).replace("_", r"\_").replace("%", r"\%").replace("&", r"\&").replace("#", r"\#")

    def ascii_of(s):
        return "".join(ch if ord(ch) < 128 else "?" for ch in str(s))

    def seeds_str(s):
        try:
            return [int(x) for x in ast.literal_eval(str(s))]
        except (ValueError, SyntaxError):
            return []

    def tb(x):
        """只把真正的 True 当 True(混进 NaN 的 object 列里 bool(nan) 会是 True)。"""
        return isinstance(x, (bool, np.bool_)) and bool(x)

    def combine(d, lo, hi, rms):
        """合成区间：test 抽样(由 bootstrap CI 反推)+ 训练随机性(零分布 rms)，正态近似。"""
        se = (hi - lo) / 3.92 if np.isfinite(lo) and np.isfinite(hi) else np.nan
        if not (np.isfinite(se) and np.isfinite(rms)):
            return np.nan, np.nan, np.nan
        sd = float(np.sqrt(se ** 2 + rms ** 2))
        return d - 1.96 * sd, d + 1.96 * sd, sd

    def tex_table(caption, label, header, rows, colspec, notes="", wide=None, sink=None):
        """booktabs 表；tabular 用 \resizebox 缩到栏宽(需要 graphicx)，列多时用跨栏 table*。"""
        wide = len(header) > 7 if wide is None else bool(wide)
        env = "table*" if wide else "table"
        out = [rf"\begin{{{env}}}[t]", r"\centering", rf"\caption{{{caption}}}", rf"\label{{{label}}}",
               r"\setlength{\tabcolsep}{3pt}", r"\resizebox{\linewidth}{!}{%",
               rf"\begin{{tabular}}{{{colspec}}}", r"\toprule", " & ".join(header) + r" \\", r"\midrule"]
        for r_ in rows:  # 规则线原样；已经以 \\ 结尾的原样；其余补 \\
            out.append(r_ if (r_.strip() in ("\\midrule", "\\addlinespace") or r_.rstrip().endswith("\\\\"))
                       else (r_ + r" \\"))
        out += [r"\bottomrule", r"\end{tabular}}"]
        if notes:
            out.append(rf"\par\smallskip\parbox{{0.97\linewidth}}{{\scriptsize {notes}}}")
        out.append(rf"\end{{{env}}}")
        (tex if sink is None else sink).append("\n".join(out) + "\n")

    say("=" * 100)
    say(f"论文表格 + 训练随机性噪声底(27 号 v6.2，第16批c 2026-10-02g)  {time.strftime('%Y-%m-%d %H:%M:%S')}  主模型={main_run}/{main_variant}")
    say("=" * 100)

    # ------------------------------------------------------------------ [0] 输入盘点 + 数据规模
    say("\n[0] 输入盘点(来源目录 / 修改时间)")
    src_list = [("主模型导出", os.path.join(results_root, main_run, "predictions_test.parquet")),
                ("第9批 19号(主结果、v8对v3)", os.path.join(compare_main, "headline_all_seeds.csv")),
                ("第9批 25号(Head A 1108基因)", os.path.join(head_a_main, "metrics_all_seeds.csv")),
                ("26号 基线", os.path.join(baselines, "table2_head_bc.csv"))]
    src_list += [(f"消融 19号 {os.path.basename(d)}", os.path.join(d, "delta_vs_reference.csv")) for d in ablation_compare]
    src_list += [(f"消融 25号 {os.path.basename(d)}", os.path.join(d, "delta_vs_reference.csv")) for d in ablation_head_a]
    for nm, p in src_list:
        say(f"  {nm:34s} {mtime(p):18s} {p}")
    pt_path = os.path.join(results_root, main_run, "predictions_test.parquet")
    pv_path = os.path.join(results_root, main_run, "predictions_val.parquet")
    if not (os.path.exists(pt_path) and os.path.exists(pv_path)):
        raise SystemExit(f"找不到主模型导出 {pt_path} / {pv_path}(先 ./run_all.sh --export)")
    pt, pv = pd.read_parquet(pt_path), pd.read_parquet(pv_path)
    facts = {}
    for sp_, df_ in (("val", pv), ("test", pt)):
        yb_ = pd.to_numeric(df_["y_b_true"], errors="coerce").to_numpy(np.float64)
        yc_ = df_["y_c_true"].astype(str).to_numpy()
        facts[sp_] = dict(rows=len(df_), genes=int(df_["gene_id"].nunique()), tfs=int(df_["tf_depleted"].nunique()),
                          sig=int(np.isfinite(yb_).sum()), down=int((yc_ == "down").sum()), up=int((yc_ == "up").sum()),
                          d_in_lg=float(df_["D_in_Lg"].astype(bool).mean()) if "D_in_Lg" in df_.columns else np.nan)
        say(f"  {sp_}: {facts[sp_]['rows']} 对(基因×被耗竭TF)，{facts[sp_]['genes']} 个基因 × {facts[sp_]['tfs']} 个 TF；显著 "
            f"{facts[sp_]['sig']}(down {facts[sp_]['down']} / up {facts[sp_]['up']})；D∈L_g 占 {facts[sp_]['d_in_lg']:.1%}")

    # ------------------------------------------------------------------ [1] 训练随机性零分布
    say("\n[1] 训练随机性零分布(同配置 v8 的 k-seed 集成两两相减；k=1：10 对单 seed；k=2：15 组互不重叠的 2+2；同一 test 集、不重采样)")
    yv = pv["y_c_true"].map(cls).to_numpy().astype(np.int64)
    yt = pt["y_c_true"].map(cls).to_numpy().astype(np.int64)
    yb_t = pd.to_numeric(pt["y_b_true"], errors="coerce").to_numpy(np.float64)
    ya_t = pt["y_a_true"].to_numpy(np.float64)
    genes_t = pt["gene_id"].to_numpy()
    tf_code = pd.factorize(pt["tf_depleted"])[0]
    n_tf = int(tf_code.max()) + 1
    _, first_gene = np.unique(genes_t, return_index=True)
    in_lg = pt["D_in_Lg"].to_numpy().astype(bool) if "D_in_Lg" in pt.columns else None
    layers = [("D∈L_g", in_lg), ("D∉L_g", ~in_lg)] if in_lg is not None else []

    def seeds_of(df_, suf):
        out = []
        for c in df_.columns:
            m_ = re.match(r"^y_b_seed(\d+)(_ph)?$", str(c))
            if m_ and (m_.group(2) or "") == suf:
                out.append(int(m_.group(1)))
        return sorted(out)

    def probs_of(df_, s, suf):
        pdn = df_[f"p_down_seed{s}{suf}"].to_numpy(np.float64)
        pup = df_[f"p_up_seed{s}{suf}"].to_numpy(np.float64)
        return np.column_stack([pdn, np.clip(1.0 - pdn - pup, 1e-12, 1.0), pup])

    def metrics(ya, yb, pr_t, pred_off, idx=None):
        """跟 19 号 metrics 同定义；idx=None 是 test 全部(含 A)，否则是分层的行下标(只算 B/C，跟 19 号 [6] 一样)。"""
        out = {}
        rows = slice(None) if idx is None else idx
        if idx is None and ya is not None:
            out["A_r_gene"] = pearson(ya[first_gene], ya_t[first_gene])
        ybt, ybp, tfc = yb_t[rows], yb[rows], tf_code[rows]
        sg = np.isfinite(ybt)
        p_, t_, c_ = ybp[sg], ybt[sg], tfc[sg]
        out["B_r"] = pearson(p_, t_)
        out["B_sign_acc"] = float(np.mean(np.sign(p_) == np.sign(t_))) if len(p_) else float("nan")
        cnt = np.maximum(np.bincount(c_, minlength=n_tf), 1)
        mp = np.bincount(c_, weights=p_, minlength=n_tf) / cnt
        mt = np.bincount(c_, weights=t_, minlength=n_tf) / cnt
        out["B_r_within_tf"] = pearson(p_ - mp[c_], t_ - mt[c_])
        yct, prt = yt[rows], pr_t[rows]
        out["C_auroc_down"], out["C_auprc_down"] = auroc_auprc(prt[:, 0], yct == 0)
        out["C_auroc_up"], out["C_auprc_up"] = auroc_auprc(prt[:, 2], yct == 2)
        out["C_macro_f1_offset"] = macro_f1(pred_off[rows], yct)
        return out

    null_rows = []
    t1 = time.time()
    for var, suf in (("perhead", "_ph"), ("selected", "")):
        ss = seeds_of(pt, suf)
        if len(ss) < 4:
            say(f"  {var}: 只有 {len(ss)} 个 seed，造不出 2 对 2 的零分布，跳过")
            continue
        cache = {}

        def ens_m(sub):
            if sub not in cache:
                ya = np.mean([pt[f"y_a_seed{s}{suf}"].to_numpy(np.float64) for s in sub], 0)
                yb = np.mean([pt[f"y_b_seed{s}{suf}"].to_numpy(np.float64) for s in sub], 0)
                pv_e = np.mean([probs_of(pv, s, suf) for s in sub], 0)
                pt_e = np.mean([probs_of(pt, s, suf) for s in sub], 0)
                off = apply_off(pt_e, tune(pv_e, yv))
                res = {("test", k): v for k, v in metrics(ya, yb, pt_e, off).items()}
                for lname, mask in layers:
                    res.update({(lname, k): v for k, v in metrics(ya, yb, pt_e, off, np.flatnonzero(mask)).items()})
                cache[sub] = res
            return cache[sub]

        for k in (1, 2):
            subs = list(itertools.combinations(ss, k))
            pairs = [(a, b) for a in subs for b in subs if a < b and not set(a) & set(b)]
            deltas = {}
            for a, b in pairs:
                ma, mb = ens_m(a), ens_m(b)
                for key in ma:
                    deltas.setdefault(key, []).append(ma[key] - mb[key])
            for (scope, met), arr in deltas.items():
                arr = np.asarray(arr, np.float64)
                arr = arr[np.isfinite(arr)]
                null_rows.append(dict(variant=var, scope=scope, metric=met, k=k, n_pairs=len(arr), seeds=str(ss),
                                      rms=float(np.sqrt(np.mean(arr ** 2))) if len(arr) else np.nan,
                                      max_abs=float(np.max(np.abs(arr))) if len(arr) else np.nan))
        say(f"  {var}: seed {ss}，单 seed 集成 {sum(len(x) == 1 for x in cache)} 个、2-seed 集成 "
            f"{sum(len(x) == 2 for x in cache)} 个已算(累计 {time.time() - t1:.0f} 秒)")
    # Head A 全基因(25 号逐基因 csv 里 v8 的逐 seed 列)
    pgA, pgA_path = None, None
    for p in head_a_seed_csv:
        q = rd(p)
        if q is not None and sum(bool(re.match(rf"^{re.escape(main_run)}_seed\d+$", c)) for c in q.columns) >= 4:
            pgA, pgA_path = q, p
            break
    strataA = {}
    if pgA is None:
        say(f"  Head A 全基因：{head_a_seed_csv} 里都没有 {main_run} ≥4 个 seed 的列，跳过")
    else:
        colsA = sorted(c for c in pgA.columns if re.match(rf"^{re.escape(main_run)}_seed\d+$", c))
        ssA = tuple(int(c.rsplit("seed", 1)[1]) for c in colsA)
        yA = pd.to_numeric(pgA["y_a_true"], errors="coerce").to_numpy(np.float64)
        spA = pgA["split"].astype(str).to_numpy()
        ig = pgA["in_grid"].astype(str).str.lower().isin(["true", "1"]).to_numpy()
        hs = pgA["has_sites"].astype(str).str.lower().isin(["true", "1"]).to_numpy()
        strataA = {"全部": np.ones(len(pgA), bool), "网格": ig, "网格外": ~ig, "网格外有位点": ~ig & hs,
                   "网格外无位点": ~ig & ~hs}
        VA = {s: pgA[f"{main_run}_seed{s}"].to_numpy(np.float64) for s in ssA}
        for k in (1, 2):
            subs = list(itertools.combinations(ssA, k))
            pairs = [(a, b) for a in subs for b in subs if a < b and not set(a) & set(b)]
            for split_ in ("test", "val"):
                for nm, msk in strataA.items():
                    ix = (spA == split_) & msk & np.isfinite(yA)
                    if ix.sum() < 5:
                        continue
                    rr = {sub: pearson(np.mean([VA[s] for s in sub], 0)[ix], yA[ix]) for sub in subs}
                    arr = np.asarray([rr[a] - rr[b] for a, b in pairs], np.float64)
                    arr = arr[np.isfinite(arr)]
                    null_rows.append(dict(variant="perhead", scope=f"A全基因:{split_}:{nm}", metric="A_r", k=k,
                                          n_pairs=len(arr), seeds=str(list(ssA)),
                                          rms=float(np.sqrt(np.mean(arr ** 2))) if len(arr) else np.nan,
                                          max_abs=float(np.max(np.abs(arr))) if len(arr) else np.nan))
        say(f"  Head A 全基因：读 {pgA_path}，{main_run} 的 seed {list(ssA)}")
    null_df = pd.DataFrame(null_rows)
    null_df.to_csv(os.path.join(outdir, "null_seed_variability.csv"), index=False)

    def null_rms(scope, metric, k, var="perhead"):
        """返回 (rms, 来源说明)。k=1/2 实测；其它 k 按 rms_2·sqrt(2/k) 外推。"""
        if null_df.empty:
            return np.nan, "无零分布"
        q = null_df[(null_df["variant"] == var) & (null_df["scope"] == scope) & (null_df["metric"] == metric)]
        if k in (1, 2):
            v = q[q["k"] == k]["rms"]
            return (float(v.iloc[0]), f"实测k={k}") if len(v) else (np.nan, "无")
        v = q[q["k"] == 2]["rms"]
        return (float(v.iloc[0]) * np.sqrt(2.0 / max(int(k), 1)), f"外推k={k}") if len(v) else (np.nan, "无")

    def null_rms_cons(scope, metric, k, var="perhead"):
        """2026-10-02a 第12批：保守版。k≥3 时不外推，直接用实测 k=2 的 rms；k=1/2 跟 null_rms 相同。"""
        if int(k) in (1, 2):
            return null_rms(scope, metric, int(k), var)
        if null_df.empty:
            return np.nan, "无零分布"
        q = null_df[(null_df["variant"] == var) & (null_df["scope"] == scope) & (null_df["metric"] == metric)
                    & (null_df["k"] == 2)]["rms"]
        return (float(q.iloc[0]), f"保守k={k}(用实测k=2)") if len(q) else (np.nan, "无")

    def sign_p(b, k):
        """2026-10-02a 第12批：逐 seed 符号检验，双侧精确二项 p。b=变好的 seed 数。"""
        try:
            b, k = int(b), int(k)
        except (TypeError, ValueError):
            return np.nan
        if k <= 0 or b < 0 or b > k:
            return np.nan
        tail = sum(math.comb(k, i) for i in range(min(b, k - b) + 1)) / 2.0 ** k
        return float(min(1.0, 2.0 * tail))

    if not null_df.empty:
        show = ["A_r_gene", "B_r", "B_r_within_tf", "B_sign_acc", "C_auroc_down", "C_auprc_down", "C_auprc_up",
                "C_macro_f1_offset"]
        say("  均方根 rms(最大|Δ|)，perhead：")
        for scope in ["test"] + [l_ for l_, _ in layers]:
            cells = []
            for met in show:
                r1, _ = null_rms(scope, met, 1)
                r2, _ = null_rms(scope, met, 2)
                q2 = null_df[(null_df["variant"] == "perhead") & (null_df["scope"] == scope) & (null_df["metric"] == met)
                             & (null_df["k"] == 2)]
                if np.isfinite(r2):
                    cells.append(f"{SHORT[met]} {r2:.3f}({q2['max_abs'].iloc[0]:.3f}) k1 {r1:.3f}")
            if cells:
                say(f"    {scope:6s} k=2: " + "  ".join(cells))
        if strataA:
            for split_ in ("test", "val"):
                cells = []
                for nm in strataA:
                    r1, _ = null_rms(f"A全基因:{split_}:{nm}", "A_r", 1)
                    r2, _ = null_rms(f"A全基因:{split_}:{nm}", "A_r", 2)
                    if np.isfinite(r2):
                        cells.append(f"{nm} {r2:.3f}(k1 {r1:.3f})")
                say(f"    Head A 全基因 {split_} k=2: " + "  ".join(cells))
        rat = []
        for met in show:
            r1, _ = null_rms("test", met, 1)
            r2, _ = null_rms("test", met, 2)
            if np.isfinite(r1) and r1 > 0 and np.isfinite(r2):
                rat.append(r2 / r1)
        if rat:
            say(f"  rms_2/rms_1 实测中位 {np.median(rat):.2f}(范围 {min(rat):.2f}~{max(rat):.2f}；方差按 1/k 缩的理论值 0.71。"
                "偏离很大时 k≥3 的外推别信)")
        say("  读法：两个 2-seed 集成之间、只因换 seed 就能差出 rms 这么多；比较 2 对 2 的消融时，|Δ| 至少要明显大于 ~2×rms 才算\n"
            "  超出训练随机性。下面 [4][5] 的\"合成区间\"把它和 test 抽样噪声合在一起。")

    # ------------------------------------------------------------------ [2] Table 1 主结果
    say(f"\n[2] Table 1 主结果({main_run}，5-seed 集成；来源 {compare_main}/headline_all_seeds.csv 等)")
    hl = rd(os.path.join(compare_main, "headline_all_seeds.csv"))
    mA = rd(os.path.join(head_a_main, "metrics_all_seeds.csv"))
    t1b = rd(os.path.join(baselines, "table1_head_a.csv"))
    t2b = rd(os.path.join(baselines, "table2_head_bc.csv"))
    t3b = rd(os.path.join(baselines, "table3_within_tf_auc.csv"))
    t4b = rd(os.path.join(baselines, "table4_strata.csv"))
    t5b = rd(os.path.join(baselines, "table5_stack.csv"))
    dmb = rd(os.path.join(baselines, "delta_model_vs_baseline.csv"))
    t6 = rd(os.path.join(baselines, "table6_fair_ensemble.csv"))   # 2026-10-02c 第14批：26 号 v5 [2d]
    t6v = rd(os.path.join(baselines, "table6_values.csv"))
    t6n = rd(os.path.join(baselines, "table6_null.csv"))          # 2026-10-02d 第15批：26 号 v6 的零分布
    tb7 = rd(os.path.join(baselines, "table_b7_tuning.csv"))      # 2026-10-02d 第15批：26 号 v6 的 B7 选轮记录
    mdl26 = None
    if t2b is not None:
        cand = [m_ for m_ in t2b["model"].astype(str).unique() if m_.startswith("模型")]
        mdl26 = cand[0] if cand else None
    mdlA26 = None
    if t1b is not None:
        cand = [m_ for m_ in t1b["model"].astype(str).unique() if m_.startswith("模型")]
        mdlA26 = cand[0] if cand else None
    rows1 = []
    if hl is not None:
        for var in ("perhead", "selected", "perhead+stack", "selected+stack"):
            q = hl[(hl["run"] == main_run) & (hl["variant"] == var)]
            for r_ in q.itertuples():
                rows1.append(dict(block="B/C(test 全部行)", variant=var, metric=r_.metric, value=r_.ensemble,
                                  ci_lo=r_.ci_lo, ci_hi=r_.ci_hi, seed_mean=r_.seed_mean, seed_std=r_.seed_std,
                                  n_seeds=r_.n_seeds, source=os.path.join(compare_main, "headline_all_seeds.csv")))
    else:
        say(f"  ⚠ 没有 {compare_main}/headline_all_seeds.csv(第9批 19 号产出)")
    if mA is not None:
        q = mA[(mA["run"] == main_run) & (mA["split"] == "test")]
        for r_ in q.itertuples():
            rows1.append(dict(block="Head A(1108 基因分层)", variant="perhead", metric=f"A_r:{r_.stratum}", value=r_.r,
                              ci_lo=r_.ci_lo, ci_hi=r_.ci_hi, seed_mean=r_.r_seed_mean, seed_std=r_.r_seed_std,
                              n_seeds=r_.seeds, n_genes=r_.n_genes, source=os.path.join(head_a_main, "metrics_all_seeds.csv")))
    if t1b is not None and mdlA26:
        q = t1b[(t1b["model"] == mdlA26) & (t1b["gene_set"] == "test 全部(组内中心化)")]
        for r_ in q.itertuples():
            rows1.append(dict(block="Head A(1108 基因分层)", variant="perhead", metric="A_r:组内中心化", value=r_.r,
                              ci_lo=r_.ci_lo, ci_hi=r_.ci_hi, n_genes=r_.n_genes,
                              source=os.path.join(baselines, "table1_head_a.csv")))
    if t3b is not None and mdl26:
        q = t3b[t3b["model"] == mdl26]
        for r_ in q.itertuples():
            rows1.append(dict(block="Head C TF 内 AUROC", variant="perhead", metric=f"wTF_pooled_{r_.direction}",
                              value=r_.pooled_auc, ci_lo=r_.pooled_ci_lo, ci_hi=r_.pooled_ci_hi, n_tf=r_.n_tf,
                              source=os.path.join(baselines, "table3_within_tf_auc.csv")))
            rows1.append(dict(block="Head C TF 内 AUROC", variant="perhead", metric=f"wTF_median_{r_.direction}",
                              value=r_.median_auc, ci_lo=r_.median_ci_lo, ci_hi=r_.median_ci_hi, n_tf=r_.n_tf,
                              source=os.path.join(baselines, "table3_within_tf_auc.csv")))
    T1 = pd.DataFrame(rows1)
    T1.to_csv(os.path.join(outdir, "table1_main.csv"), index=False)
    if not T1.empty:
        for blk, q in T1.groupby("block", sort=False):
            say(f"  {blk}：")
            for r_ in q.itertuples():
                if blk.startswith("B/C") and r_.variant not in ("perhead", "perhead+stack"):
                    continue
                extra = (f"  逐seed {r_.seed_mean:.3f}±{r_.seed_std:.3f}" if "seed_mean" in T1.columns
                         and np.isfinite(getattr(r_, "seed_mean", np.nan)) else "")
                say(f"    {r_.variant:14s} {SHORT.get(r_.metric, r_.metric):26s} {f3(r_.value)}[{f3(r_.ci_lo)},{f3(r_.ci_hi)}]"
                    + extra)

    def t1v(metric, variant="perhead"):
        if T1.empty:
            return np.nan, np.nan, np.nan
        q = T1[(T1["metric"] == metric) & (T1["variant"] == variant)]
        return (float(q["value"].iloc[0]), float(q["ci_lo"].iloc[0]), float(q["ci_hi"].iloc[0])) if len(q) \
            else (np.nan, np.nan, np.nan)

    # ------------------------------------------------------------------ [3] Table 2/3 对基线
    say(f"\n[3] Table 2(Head B/C 对基线)、Table 3(Head A 对基线)；来源 {baselines}/table*.csv(26 号)")
    T2 = pd.DataFrame()
    stack_name = None
    if t2b is not None and mdl26:
        piv = t2b.pivot_table(index="model", columns="metric", values="value", aggfunc="first")
        lo_ = t2b.pivot_table(index="model", columns="metric", values="ci_lo", aggfunc="first")
        hi_ = t2b.pivot_table(index="model", columns="metric", values="ci_hi", aggfunc="first")
        order = [m_ for m_ in NAME if m_ in piv.index and m_.startswith("B")] + \
            [m_ for m_ in piv.index if m_.startswith("ORACLE")] + [mdl26]
        rows2 = []
        for m_ in order:
            row = dict(model=m_, label=NAME.get(m_, ("TF-Layout (ours)",))[0] if m_ != mdl26 else "TF-Layout (ours)")
            for k26, kl in B26.items():
                row[kl] = piv.loc[m_, k26] if k26 in piv.columns else np.nan
                row[kl + "_lo"] = lo_.loc[m_, k26] if k26 in lo_.columns else np.nan
                row[kl + "_hi"] = hi_.loc[m_, k26] if k26 in hi_.columns else np.nan
            if t3b is not None:
                for d_ in ("dn", "up"):
                    q = t3b[(t3b["model"] == m_) & (t3b["direction"] == d_)]
                    row[f"wtf_{d_}"] = float(q["pooled_auc"].iloc[0]) if len(q) else np.nan
                    row[f"wtf_{d_}_lo"] = float(q["pooled_ci_lo"].iloc[0]) if len(q) else np.nan
                    row[f"wtf_{d_}_hi"] = float(q["pooled_ci_hi"].iloc[0]) if len(q) else np.nan
            rows2.append(row)
        if t5b is not None and len(t5b):
            stack_name = str(t5b["stack"].iloc[0])
            for stk_ in list(dict.fromkeys(t5b["stack"].astype(str))):   # 2026-10-02c 第14批：每个堆叠一行(B6、B7a_bag、B7b_bag)
                qs_ = t5b[t5b["stack"].astype(str) == stk_]
                refs_ = [x_ for x_ in dict.fromkeys(qs_["reference"].astype(str)) if x_ != mdl26]
                base_short = NAME.get(refs_[0], (refs_[0],))[0].split()[0] if refs_ else stk_
                row = dict(model=stk_, label=f"TF-Layout + {base_short} (stacked)")
                for r_ in qs_.drop_duplicates("metric").itertuples():
                    row[B26.get(r_.metric, r_.metric)] = r_.stack_value
                rows2.append(row)
        T2 = pd.DataFrame(rows2)
        T2.to_csv(os.path.join(outdir, "table2_head_bc.csv"), index=False)
        cols2 = ["B_r", "B_sign_acc", "B_r_within_tf", "C_auroc_down", "C_auroc_up", "C_auprc_down", "C_auprc_up",
                 "C_macro_f1_offset", "wtf_dn", "wtf_up"]
        say("    " + "".ljust(28) + "".join(f"{SHORT.get(c, c):>9}" for c in cols2))
        for r_ in rows2:
            say("    " + ascii_of(r_["label"])[:27].ljust(28) + "".join(f"{f3(r_.get(c, np.nan)):>9}" for c in cols2))
        # LaTeX：粗体=除 Oracle/堆叠外每列最好；我们那行下面一行给 CI
        comp = [r_ for r_ in rows2 if not str(r_["model"]).startswith(("ORACLE", "STACK"))]
        best = {c: np.nanmax([r_.get(c, np.nan) for r_ in comp]) for c in cols2}
        trows = []
        for r_ in rows2:
            lab = NAME.get(r_["model"], (None, r"\textbf{TF-Layout (ours)}"))[1] if r_["model"] != mdl26 \
                else r"\textbf{TF-Layout (ours)}"
            if str(r_["model"]).startswith("STACK"):
                lab = r"\textit{" + str(r_["label"]) + "}"
                if not any(str(x_).startswith(r"\textit{TF-Layout +") for x_ in [t_.split(" & ")[0] for t_ in trows if isinstance(t_, str)]):
                    trows.append(r"\midrule")
            cells = []
            for c in cols2:
                v = r_.get(c, np.nan)
                s_ = f3(v) if np.isfinite(v) else "--"
                if np.isfinite(v) and not str(r_["model"]).startswith(("ORACLE", "STACK")) and abs(v - best[c]) < 5e-4:
                    s_ = rf"\textbf{{{s_}}}"
                cells.append(s_)
            trows.append(lab + " & " + " & ".join(cells))
            if r_["model"] == mdl26:
                trows.append(r"\quad {\scriptsize 95\% CI} & " + " & ".join(
                    (rf"{{\scriptsize [{f3(r_.get(c + '_lo'))},{f3(r_.get(c + '_hi'))}]}}"
                     if np.isfinite(r_.get(c + "_lo", np.nan)) else "") for c in cols2))
        tex_table(r"Head B/C on held-out chromosomes (test: %d gene$\times$TF pairs, %d genes, %d significant). "
                  r"Bold: best non-oracle method. wTF-AUROC: pooled within-TF AUROC (TFs with $\geq$5 positives)."
                  % (facts["test"]["rows"], facts["test"]["genes"], facts["test"]["sig"]),
                  "tab:headbc", ["Method"] + [TEXM.get(c, c) for c in cols2], trows, "l" + "c" * len(cols2),
                  notes="Baselines are fit on training chromosomes; ridge/logistic hyper-parameters and Head C decision "
                        "offsets tuned on validation chromosomes (26\\_paper\\_baselines.py). Oracle uses test labels of "
                        "the same gene under other TFs (decomposition reference only). Stacking is fit on validation.")
    else:
        say("  ⚠ 没有 26 号 table2_head_bc.csv，跳过 Table 2")
    T3 = pd.DataFrame()
    if t1b is not None and mdlA26:
        sets = list(dict.fromkeys(t1b["gene_set"].astype(str)))
        rows3 = []
        order = [m_ for m_ in NAME if m_ in set(t1b["model"])] + [mdlA26]
        for m_ in order:
            q = t1b[t1b["model"] == m_].set_index("gene_set")
            row = dict(model=m_, label="TF-Layout (ours)" if m_ == mdlA26 else NAME[m_][0])
            for s_ in sets:
                if s_ in q.index:
                    row[s_] = q.loc[s_, "r"]
                    row[s_ + "_lo"], row[s_ + "_hi"] = q.loc[s_, "ci_lo"], q.loc[s_, "ci_hi"]
                    row[s_ + "_n"] = q.loc[s_, "n_genes"]
                    row[s_ + "_d"] = q.loc[s_, "delta_model_minus_base"]
                    row[s_ + "_dlo"], row[s_ + "_dhi"] = q.loc[s_, "delta_ci_lo"], q.loc[s_, "delta_ci_hi"]
            rows3.append(row)
        T3 = pd.DataFrame(rows3)
        T3.to_csv(os.path.join(outdir, "table3_head_a.csv"), index=False)
        gs_en = {"test 全部": "All", "test 网格": "Grid", "test 网格外": "Off-grid", "test 网格外有位点": "Off+sites",
                 "test 网格外无位点": "Off,no-sites", "test 全部(组内中心化)": "All,centered"}  # 2026-10-02a：原来打印成 ??
        say("    " + "".ljust(26) + "".join(f"{gs_en.get(s_, ascii_of(s_.replace('test ', '')))[:14]:>16}" for s_ in sets))
        for r_ in rows3:
            say("    " + ascii_of(r_["label"])[:25].ljust(26) + "".join(f"{f3(r_.get(s_, np.nan)):>16}" for s_ in sets))
        gs_tex = {"test 全部": "All", "test 网格": "Grid", "test 网格外": "Off-grid", "test 网格外有位点": "Off-grid+sites",
                  "test 网格外无位点": "Off-grid, no sites", "test 全部(组内中心化)": "All, group-centered"}
        comp = [r_ for r_ in rows3]
        best = {s_: np.nanmax([r_.get(s_, np.nan) for r_ in comp]) for s_ in sets}
        trows = []
        for r_ in rows3:
            lab = r"\textbf{TF-Layout (ours)}" if r_["model"] == mdlA26 else NAME[r_["model"]][1]
            cells = []
            for s_ in sets:
                v = r_.get(s_, np.nan)
                c_ = f3(v) if np.isfinite(v) else "--"
                if np.isfinite(v) and abs(v - best[s_]) < 5e-4:
                    c_ = rf"\textbf{{{c_}}}"
                if r_["model"] == mdlA26 and np.isfinite(r_.get(s_ + "_lo", np.nan)):
                    c_ += rf" {{\scriptsize [{f3(r_[s_ + '_lo'])},{f3(r_[s_ + '_hi'])}]}}"
                cells.append(c_)
            trows.append(lab + " & " + " & ".join(cells))
        hdr = ["Method"] + [f"{gs_tex.get(s_, texesc(s_))} ({int(rows3[-1].get(s_ + '_n', 0))})" for s_ in sets]
        tex_table(r"Head A (wild-type expression, z-scored $\log$ TPM): Pearson $r$ on held-out-chromosome genes. "
                  r"Group-centered: each gene group (grid / off-grid with sites / off-grid without sites) centered "
                  r"separately before pooling, removing between-group mean differences.",
                  "tab:heada", hdr, trows, "l" + "c" * len(sets), wide=True)
    else:
        say("  ⚠ 没有 26 号 table1_head_a.csv，跳过 Table 3")

    # ------------------------------------------------------------------ [4] Table 4 消融
    say(f"\n[4] Table 4 消融(在 {main_run} 上；参照 {main_run}/{main_variant}；自动取 seed 最多的来源)。"
        "†=bootstrap CI 不含0(只含 test 抽样)；‡=合成区间不含0(再加训练随机性)；§=保守合成区间不含0(k≥3 时不外推，只标在 k≥3 的行)")
    cmp_cache = {}

    def load_cmp(d):
        if d not in cmp_cache:
            cmp_cache[d] = dict(delta=rd(os.path.join(d, "delta_vs_reference.csv")),
                                pair=rd(os.path.join(d, "pairwise_delta.csv")),
                                strata=rd(os.path.join(d, "strata_delta.csv")),
                                facts=rjson(os.path.join(d, "compare_facts.json")))
        return cmp_cache[d]

    def find_delta(run, ref_run, var, dirs):
        """在 dirs 里找 (ref_run/var → run/var) 的差值行，返回 (df, k, 来源)；同一实验取 seed 最多的。"""
        cands = []
        for d in dirs:
            C = load_cmp(d)
            ref_name = f"{ref_run}/{var}"
            if C["delta"] is not None and C["facts"] is not None:
                q = C["delta"][(C["delta"]["reference"] == ref_name) & (C["delta"]["run"] == run)
                               & (C["delta"]["variant"] == var)]
                if len(q):
                    cands.append((len(C["facts"].get("seeds", [])), -len(cands), q,
                                  f"{d}/delta_vs_reference.csv(seed {C['facts'].get('seeds')})"))
            if C["pair"] is not None:
                q = C["pair"][(C["pair"]["reference"] == ref_name) & (C["pair"]["run"] == run)
                              & (C["pair"]["variant"] == var)]
                if len(q):
                    ss_ = seeds_str(q["seeds"].iloc[0])
                    cands.append((len(ss_), -len(cands), q, f"{d}/pairwise_delta.csv(seed {ss_})"))
        if not cands:
            return None, 0, ""
        k, _, q, src = max(cands, key=lambda x: (x[0], x[1]))
        return q, k, src

    def find_head_a(run, ref_run, dirs):
        cands = []
        for d in dirs:
            q0 = rd(os.path.join(d, "delta_vs_reference.csv"))
            if q0 is None:
                continue
            q = q0[(q0["run"] == run) & (q0["reference"] == ref_run)]
            if len(q):
                k = len(seeds_str(q["seeds"].iloc[0]))
                cands.append((k, -len(cands), q, f"{d}/delta_vs_reference.csv(seed {seeds_str(q['seeds'].iloc[0])})"))
        if not cands:
            return None, 0, ""
        k, _, q, src = max(cands, key=lambda x: (x[0], x[1]))
        return q, k, src

    def delta_rows(q, k, src, label, run, scope_prefix="", var="perhead"):
        out = []
        for r_ in q.itertuples():
            met = r_.metric
            if met not in M_LONG:
                continue
            rms, how = null_rms("test", met, k, var)
            lo_t, hi_t, sd_t = combine(r_.delta, r_.ci_lo, r_.ci_hi, rms)
            rms_c, _ = null_rms_cons("test", met, k, var)
            lo_c, hi_c, _ = combine(r_.delta, r_.ci_lo, r_.ci_hi, rms_c)
            sb_ = getattr(r_, "n_seeds_better", np.nan)
            out.append(dict(label=label, run=run, scope=scope_prefix + "test", metric=met, k=k, ref=r_.ref,
                            value=r_.value, delta=r_.delta, ci_lo=r_.ci_lo, ci_hi=r_.ci_hi,
                            seed_delta=getattr(r_, "paired_seed_delta_mean", np.nan),
                            seeds_better=getattr(r_, "n_seeds_better", np.nan),
                            null_rms=rms, null_how=how, tot_lo=lo_t, tot_hi=hi_t, tot_sd=sd_t,
                            test_sig=bool(r_.ci_lo > 0 or r_.ci_hi < 0),
                            tot_sig=bool(np.isfinite(lo_t) and (lo_t > 0 or hi_t < 0)),
                            null_rms_cons=rms_c, cons_lo=lo_c, cons_hi=hi_c,
                            cons_sig=bool(np.isfinite(lo_c) and (lo_c > 0 or hi_c < 0)), sign_p=sign_p(sb_, k), source=src))
        return out

    def head_a_rows(q, k, src, label, run, split_="test"):
        out = []
        for r_ in q[q["split"] == split_].itertuples():
            rms, how = null_rms(f"A全基因:{split_}:{r_.stratum}", "A_r", k)
            lo_t, hi_t, sd_t = combine(r_.delta, r_.ci_lo, r_.ci_hi, rms)
            rms_c, _ = null_rms_cons(f"A全基因:{split_}:{r_.stratum}", "A_r", k)
            lo_c, hi_c, _ = combine(r_.delta, r_.ci_lo, r_.ci_hi, rms_c)
            sb_ = r_.seeds_better
            out.append(dict(label=label, run=run, scope=f"A全基因:{split_}:{r_.stratum}", metric="A_r", k=k,
                            ref=r_.r_ref, value=r_.r_run, delta=r_.delta, ci_lo=r_.ci_lo, ci_hi=r_.ci_hi,
                            seed_delta=r_.seed_delta_mean, seeds_better=r_.seeds_better, null_rms=rms, null_how=how,
                            tot_lo=lo_t, tot_hi=hi_t, tot_sd=sd_t,
                            test_sig=bool(np.isfinite(r_.delta) and (r_.ci_lo > 0 or r_.ci_hi < 0)),
                            tot_sig=bool(np.isfinite(lo_t) and (lo_t > 0 or hi_t < 0)),
                            null_rms_cons=rms_c, cons_lo=lo_c, cons_hi=hi_c,
                            cons_sig=bool(np.isfinite(lo_c) and (lo_c > 0 or hi_c < 0)), sign_p=sign_p(sb_, k), source=src))
        return out

    abl_rows = []
    for run, lab_a, lab_t in ablations:
        q, k, src = find_delta(run, main_run, main_variant, ablation_compare)
        if q is None:
            say(f"  {run}: 在 {ablation_compare} 里找不到对 {main_run}/{main_variant} 的差值，跳过")
            continue
        abl_rows += delta_rows(q, k, src, lab_a, run, var=main_variant)
        qa, ka, srca = find_head_a(run, main_run, ablation_head_a)
        if qa is not None:
            abl_rows += head_a_rows(qa, ka, srca, lab_a, run, "test")
            abl_rows += head_a_rows(qa, ka, srca, lab_a, run, "val")
        say(f"  {lab_a}({run})：B/C 来自 {src}；Head A 全基因来自 {srca or '无'}")
    T4 = pd.DataFrame(abl_rows)
    T4.to_csv(os.path.join(outdir, "table4_ablation_long.csv"), index=False)
    show4 = ["A_r_gene", "B_r", "B_r_within_tf", "B_sign_acc", "C_auroc_down", "C_auroc_up", "C_auprc_down",
             "C_auprc_up", "C_macro_f1_offset"]
    if not T4.empty:
        for lab_a in dict.fromkeys(T4["label"]):
            q = T4[T4["label"] == lab_a]
            say(f"  ── {lab_a}  (k={int(q['k'].iloc[0])} seed 对 {int(q['k'].iloc[0])} seed)")
            for r_ in q[q["scope"] == "test"].set_index("metric").reindex(show4).dropna(subset=["delta"]).reset_index().itertuples():
                big_k = int(r_.k) > 2
                mk = ("†" if r_.test_sig else " ") + ("‡" if r_.tot_sig else " ") + ("§" if big_k and tb(r_.cons_sig) else " ")
                say(f"      {SHORT[r_.metric]:8s} Δ={r_.delta:+.4f}{mk} test CI({r_.ci_lo:+.3f},{r_.ci_hi:+.3f})  "
                    f"噪声 rms {r_.null_rms:.4f}[{r_.null_how}]  合成({r_.tot_lo:+.3f},{r_.tot_hi:+.3f})"
                    + (f"  保守({f3(r_.cons_lo, True)},{f3(r_.cons_hi, True)})" if big_k else "")
                    + f"  逐seed {r_.seed_delta:+.4f}({int(r_.seeds_better) if np.isfinite(r_.seeds_better) else '?'}/{int(r_.k)}"
                    + (f"，符号检验双侧p={r_.sign_p:.3f}" if np.isfinite(r_.sign_p) else "") + ")")
            qa = q[q["scope"].str.startswith("A全基因:test")]
            if len(qa):
                say("      Head A 全基因 test: " + "  ".join(
                    (f"{r_.scope.split(':')[-1]} {f3(r_.delta, True)}{'†' if r_.test_sig else ''}{'‡' if r_.tot_sig else ''}"
                     if np.isfinite(r_.delta) else f"{r_.scope.split(':')[-1]} 无定义(常数预测)")
                    for r_ in qa.itertuples()))
            qv = q[q["scope"].str.startswith("A全基因:val")]
            if len(qv):
                say("      Head A 全基因 val : " + "  ".join(
                    (f"{r_.scope.split(':')[-1]} {f3(r_.delta, True)}{'†' if r_.test_sig else ''}{'‡' if r_.tot_sig else ''}"
                     if np.isfinite(r_.delta) else f"{r_.scope.split(':')[-1]} 无定义(常数预测)")
                    for r_ in qv.itertuples()))
        cols4 = [("A_r_gene", "test"), ("A_r", "A全基因:test:全部"), ("B_r", "test"), ("B_r_within_tf", "test"),
                 ("C_auroc_down", "test"), ("C_auprc_down", "test"), ("C_auprc_up", "test"), ("C_macro_f1_offset", "test")]
        def ref_of(run_):
            qq_ = T4[T4["run"] == run_]
            out_ = []
            for met_, scope_ in cols4:
                q_ = qq_[(qq_["metric"] == met_) & (qq_["scope"] == scope_)]
                out_.append(round(float(q_["ref"].iloc[0]), 6) if len(q_) and np.isfinite(float(q_["ref"].iloc[0])) else None)
            return tuple(out_)

        groups4, order4, disp4 = {}, [], {}   # 2026-10-02f v6.1：按 (k, B_r|TF 与 C_APdn 的参照值) 分组，每组一行参照
        for run, lab_a, lab_t in ablations:
            q = T4[T4["run"] == run]
            if q.empty:
                continue
            cells = []
            for met, scope in cols4:
                r_ = q[(q["metric"] == met) & (q["scope"] == scope)]
                if not len(r_) or not np.isfinite(r_["delta"].iloc[0]):
                    cells.append("--")
                    continue
                r_ = r_.iloc[0]
                mk = (r"\dagger" if r_["test_sig"] else "") + (r"\ddagger" if r_["tot_sig"] else "") + \
                     (r"\S" if int(r_["k"]) > 2 and tb(r_.get("cons_sig", False)) else "")
                cells.append(f"${r_['delta']:+.3f}" + (f"^{{{mk}}}$" if mk else "$"))
            rf_ = ref_of(run)
            key4 = (int(q["k"].iloc[0]), rf_[3], rf_[5])   # cols4 的第 3、5 列 = B_r|TF、C_APdn；缺别的指标的实验不会另起一组
            if key4 not in groups4:
                groups4[key4] = []
                order4.append(key4)
                disp4[key4] = list(rf_)
            else:
                disp4[key4] = [a_ if a_ is not None else b_ for a_, b_ in zip(disp4[key4], rf_)]
            groups4[key4].append(f"{lab_t} ($k={int(q['k'].iloc[0])}$) & " + " & ".join(cells))
        trows = []
        for gi, key4 in enumerate(order4):
            if gi:
                trows.append(r"\midrule")
            trows.append(rf"TF-Layout (reference, $k={key4[0]}$ seeds) & "
                         + " & ".join("--" if v_ is None else f"{v_:.3f}" for v_ in disp4[key4]))
            trows.append(r"\addlinespace")
            trows += groups4[key4]
        tex_table(r"Component ablations on the final model (ensembles of $k$ seeds vs.\ the reference trained with the "
                  r"same seeds; $\Delta$ = ablation $-$ reference, test set). $\dagger$: 95\% gene-cluster bootstrap CI "
                  r"excludes 0 (test sampling only). $\ddagger$: interval that additionally includes seed-to-seed "
                  r"training variability (estimated from the reference model's 5 seeds) excludes 0. $\S$ (rows with $k>2$): "
                  r"the same interval without extrapolating the 2-seed variability to $k$ seeds (conservative) also excludes 0.",
                  "tab:ablation", ["Variant"] + [TEXM.get(m_ if s_ == "test" else "A_all", m_) for m_, s_ in cols4],
                  trows, "l" + "c" * len(cols4))
    else:
        say("  ⚠ 没有任何消融差值(先跑 19 号第10/11批)")

    # ------------------------------------------------------------------ [5] Table 5 分层
    say("\n[5] Table 5 分层(D∈L_g=被耗竭 TF 在该基因启动子上有位点)：模型对 B4/B6(26 号 [2c]) + 消融(19 号 [6])")
    rows5 = []
    if t4b is not None:
        for r_ in t4b.itertuples():
            rows5.append(dict(kind="模型−基线", stratum=r_.stratum, against=r_.baseline, metric=B26.get(r_.metric, r_.metric),
                              model_value=r_.model_value, other_value=r_.base_value, delta=r_.delta, ci_lo=r_.ci_lo,
                              ci_hi=r_.ci_hi, test_sig=bool(r_.ci_lo > 0 or r_.ci_hi < 0),
                              source=os.path.join(baselines, "table4_strata.csv")))
    for run, lab_a, lab_t in ablations:
        for d in ablation_compare:
            C = load_cmp(d)
            if C["strata"] is None or C["facts"] is None:
                continue
            q = C["strata"][(C["strata"]["run"] == run) & (C["strata"]["reference"] == f"{main_run}/{main_variant}")]
            if not len(q):
                continue
            k = len(C["facts"].get("seeds", []))
            for r_ in q.itertuples():
                rms, how = null_rms(r_.stratum, r_.metric, k)
                lo_t, hi_t, _ = combine(r_.delta, r_.ci_lo, r_.ci_hi, rms)
                rms_c, _ = null_rms_cons(r_.stratum, r_.metric, k)
                lo_c, hi_c, _ = combine(r_.delta, r_.ci_lo, r_.ci_hi, rms_c)
                rows5.append(dict(kind="消融−参照", stratum=r_.stratum, against=lab_a, metric=r_.metric,
                                  model_value=r_.ref, other_value=r_.value, delta=r_.delta, ci_lo=r_.ci_lo,
                                  ci_hi=r_.ci_hi, k=k, null_rms=rms, null_how=how, tot_lo=lo_t, tot_hi=hi_t,
                                  test_sig=bool(r_.ci_lo > 0 or r_.ci_hi < 0),
                                  tot_sig=bool(np.isfinite(lo_t) and (lo_t > 0 or hi_t < 0)),
                                  cons_lo=lo_c, cons_hi=hi_c, cons_sig=bool(np.isfinite(lo_c) and (lo_c > 0 or hi_c < 0)),
                                  n_rows=r_.n_rows, n_sig=r_.n_sig, source=f"{d}/strata_delta.csv"))
            break  # 第一个有这个实验的目录(b11 优先)
    T5 = pd.DataFrame(rows5)
    T5.to_csv(os.path.join(outdir, "table5_strata_long.csv"), index=False)
    show5 = ["B_r", "B_r_within_tf", "C_auroc_down", "C_auprc_down", "C_auprc_up"]
    if not T5.empty:
        for (kind, ag), q in T5.groupby(["kind", "against"], sort=False):
            for st in ("D∈L_g", "D∉L_g"):
                qq = q[q["stratum"] == st].set_index("metric")
                cells = []
                for m_ in show5:
                    if m_ in qq.index:
                        r_ = qq.loc[m_]
                        mk = ("†" if tb(r_["test_sig"]) else "") + ("‡" if tb(r_.get("tot_sig", False)) else "") + \
                             ("§" if tb(r_.get("cons_sig", False)) and float(r_.get("k", 0) or 0) > 2 else "")
                        cells.append(f"{SHORT[m_]} {f3(r_['delta'], True)}{mk}")
                if cells:
                    say(f"  {kind:6s} {str(ag)[:22]:22s} {st}: " + "  ".join(cells))
        trows = []
        for st in ("D∈L_g", "D∉L_g"):
            q = T5[T5["stratum"] == st]
            nrow = q["n_rows"].dropna().iloc[0] if "n_rows" in q.columns and q["n_rows"].notna().any() else np.nan
            nsig = q["n_sig"].dropna().iloc[0] if "n_sig" in q.columns and q["n_sig"].notna().any() else np.nan
            st_tex = r"$D\in L_g$" if st == "D∈L_g" else r"$D\notin L_g$"
            trows.append(rf"\multicolumn{{6}}{{l}}{{\textit{{{st_tex}}}"
                         + (f" ({int(nrow)} pairs, {int(nsig)} significant)" if np.isfinite(nrow) else "") + r"} \\")
            for kind, ag, lab in (("模型−基线", "B4_pair+TF先验", r"Ours $-$ B4"), ("模型−基线", "B6_gbdt(同B5特征)", r"Ours $-$ B6"),
                                  ("模型−基线", "B7a_gbdt+flat(无WT)", r"Ours $-$ B7a"),
                                  ("模型−基线", "B7a_bag(无WT)", r"Ours $-$ B7a-bag"),
                                  ("消融−参照", "w/o site deletion", r"w/o site deletion $-$ ours")):
                qq = q[(q["kind"] == kind) & (q["against"] == ag)].set_index("metric")
                if qq.empty:
                    continue
                cells = []
                for m_ in show5:
                    if m_ not in qq.index:
                        cells.append("--")
                        continue
                    r_ = qq.loc[m_]
                    mk = (r"\dagger" if tb(r_["test_sig"]) else "") + (r"\ddagger" if tb(r_.get("tot_sig", False)) else "")
                    cells.append(f"${r_['delta']:+.3f}" + (f"^{{{mk}}}$" if mk else "$"))
                trows.append(r"\quad " + lab + " & " + " & ".join(cells))
        tex_table(r"Where the gains come from: differences stratified by whether the depleted TF $D$ has a mapped site in "
                  r"the gene's promoter. $\dagger$/$\ddagger$ as in Table~\ref{tab:ablation} ($\ddagger$ only for "
                  r"ablations, where seed variability applies).",
                  "tab:strata", ["Comparison"] + [TEXM[m_] for m_ in show5], trows, "l" + "c" * len(show5))

    # ------------------------------------------------------------------ [5b] Table 6 公平集成对照(2026-10-02c 第14批；2026-10-02d 第15批加 ‡/§、修 "??")
    say("\n[5b] Table 6 公平集成对照(26 号 [2d])：模型单 seed / 集成 对 B7 单模型 / bagging 集成；†=按基因整群 bootstrap CI 不含0(只含 test 抽样)；"
        "‡=再加训练随机性的合成区间不含0；§=保守区间不含0(‡/§ 要 26 号 v6 才有)")
    T6 = pd.DataFrame()
    cols6 = ["B_r|TF", "C_APdn", "C_APup", "C_AUCdn", "wTF_dn", "wTF_up"]
    F6 = {"B_r|TF": "B_r_within_tf", "C_APdn": "C_auprc_down", "C_APup": "C_auprc_up", "C_AUCdn": "C_auroc_down",
          "wTF_dn": "wtf_dn", "wTF_up": "wtf_up"}
    if t6 is None or not len(t6):
        say("  (没有 26 号的 table6_fair_ensemble.csv：先 ./run_all.sh --baselines；Table 6 跳过)")
    else:
        T6 = t6.copy()
        has_tot = "tot_sig" in T6.columns
        if not has_tot:
            say("  (table6_fair_ensemble.csv 没有 tot_*/cons_* 列——还是 26 号 v5 的产出；只标 †。先 --baselines(26 号 v6)再 --tables)")
        T6.to_csv(os.path.join(outdir, "table6_fair_ensemble.csv"), index=False)
        trows = []
        if t6v is not None and len(t6v):
            say("    取值(seed 间标准差只对单 seed 均值那行)：".ljust(34) + "".join(f"{m_:>9}" for m_ in cols6))
            for lab_ in list(dict.fromkeys(t6v["label"].astype(str))):
                qv_ = t6v[t6v["label"].astype(str) == lab_].set_index("metric")
                if lab_.startswith("模型") and "单seed" in lab_:
                    en_, txt_ = r"Ours, single seed (mean $\pm$ sd)", "Ours single-seed mean"
                elif lab_.startswith("模型"):
                    en_, txt_ = r"Ours, seed ensemble", "Ours seed ensemble"
                else:
                    en_, txt_ = NAME.get(lab_, (lab_, lab_))[1], NAME.get(lab_, (lab_, lab_))[0]
                cells, cells_txt = [], []
                for m_ in cols6:
                    if m_ not in qv_.index:
                        cells.append("--")
                        cells_txt.append(f"{'--':>9}")
                        continue
                    v_, sd_ = float(qv_.loc[m_, "value"]), qv_.loc[m_].get("seed_sd", np.nan)
                    cells.append(f"${v_:.3f}" + (rf"\pm{sd_:.3f}$" if np.isfinite(sd_) else "$"))
                    cells_txt.append(f"{v_:>9.3f}")
                say("      " + ascii_of(txt_)[:30].ljust(32) + "".join(cells_txt))
                trows.append(en_ + " & " + " & ".join(cells))
            trows.append(r"\midrule")
        LAB6 = [("(i) 集成−B7a_bag(对等集成)", r"Ours (ensemble) $-$ B7a-bag", "(i) Ours ens - B7a-bag"),
                ("(j) 单seed均值−B7a(单对单)", r"Ours (single-seed mean) $-$ B7a", "(j) Ours single - B7a"),
                ("集成−B7b_bag", r"Ours (ensemble) $-$ B7b-bag", "Ours ens - B7b-bag"),
                ("单seed均值−B7b", r"Ours (single-seed mean) $-$ B7b", "Ours single - B7b"),
                ("(k) 集成增益 B7a_bag−B7a", r"Ensembling gain, B7a (bag $-$ single)", "(k) B7a bag - single"),
                ("模型集成增益 集成−单seed均值", r"Ensembling gain, Ours (ensemble $-$ single-seed mean)", "Ours ens - single")]
        for key, lab, txtlab in LAB6:
            q = T6[T6["comparison"] == key].set_index("metric")
            if q.empty:
                continue
            cells, txt = [], []
            for m_ in cols6:
                if m_ not in q.index:
                    cells.append("--")
                    continue
                r_ = q.loc[m_]
                dag = bool(r_["ci_lo"] > 0 or r_["ci_hi"] < 0)
                ddg = bool(has_tot and tb(r_.get("tot_sig", False)))
                sec = bool(has_tot and tb(r_.get("cons_sig", False)))
                mk = (r"\dagger" if dag else "") + (r"\ddagger" if ddg else "")
                cells.append(f"${r_['delta']:+.3f}" + (f"^{{{mk}}}$" if mk else "$"))
                txt.append(f"{SHORT.get(F6[m_], m_)} {r_['delta']:+.3f}{'†' if dag else ''}{'‡' if ddg else ''}{'§' if sec else ''}")
            say("    " + txtlab.ljust(26) + "  ".join(txt))
            trows.append(lab + " & " + " & ".join(cells))
        if has_tot:
            say("    (‡/§ 的区间宽度见 paper_numbers.md 的 Table 6 节；\"Ours ens - single\"两边共用同一批 seed，不给 ‡/§)")
        tex_table(r"Matched-ensemble comparison against gradient boosting on flattened layout features (B7). Top: absolute values "
                  r"(ensemble of the model's seeds; mean $\pm$ sd over its single seeds; B7 single model and B7 bagging ensemble). "
                  r"Bottom: differences; $\dagger$: 95\% gene-cluster bootstrap CI excludes 0 (test sampling only)"
                  + (r"; $\ddagger$: the interval that additionally includes re-training variability (model seeds; GBDT bagging members) "
                     r"excludes 0" if has_tot else r"; training variability is shown by the sd") + ".",
                  "tab:fair", ["Comparison / method"] + [TEXM[F6[m_]] for m_ in cols6], trows, "l" + "c" * len(cols6),
                  notes=r"B7-bag: bagged GBDT members, each trained on a random 80\% of the training genes with the hyper-parameters "
                        r"of the single model (so the bagging ensemble is, if anything, slightly disadvantaged)."
                        + (r" Re-training variability: root-mean-square difference between independent re-trainings, measured from the "
                           r"model's 5 seeds (2 vs.\ 2, scaled to 5) and from 10 GBDT members (5 vs.\ 5)." if has_tot else ""))

    # ------------------------------------------------------------------ [6] Table S1 v8 对 v3
    say(f"\n[6] Table S1：{main_run} 对 {v3_run}(扩大 Head A 训练基因集；第9批 {compare_main}/pairwise_delta.csv)")
    rowsS1 = []
    for var in ("perhead", "selected"):
        q, k, src = find_delta(main_run, v3_run, var, (compare_main,))
        if q is None:
            continue
        rowsS1 += [dict(r_, variant=var) for r_ in delta_rows(q, k, src, f"{main_run} vs {v3_run}", main_run, var=var)]
    qa, ka, srca = find_head_a(main_run, v3_run, (head_a_main,))
    if qa is not None:
        rowsS1 += [dict(r_, variant="perhead") for r_ in head_a_rows(qa, ka, srca, f"{main_run} vs {v3_run}", main_run)]
    TS1 = pd.DataFrame(rowsS1)
    TS1.to_csv(os.path.join(outdir, "tableS1_v8_vs_v3.csv"), index=False)
    if not TS1.empty:
        for var in ("perhead", "selected"):
            q = TS1[(TS1["variant"] == var) & (TS1["scope"] == "test")]
            if len(q):
                say(f"  {var}(k={int(q['k'].iloc[0])}): " + "  ".join(
                    f"{SHORT[r_.metric]} {f3(r_.delta, True)}{'†' if r_.test_sig else ''}{'‡' if r_.tot_sig else ''}"
                    for r_ in q.itertuples()))
        q = TS1[TS1["scope"].str.startswith("A全基因:test")]
        if len(q):
            say("  Head A 全基因 test: " + "  ".join(
                f"{r_.scope.split(':')[-1]} {f3(r_.delta, True)}{'†' if r_.test_sig else ''}{'‡' if r_.tot_sig else ''}"
                for r_ in q.itertuples()))
        say("  (k=5 的噪声 rms 是从 k=2 外推的，见 [1] 末尾的比值)")
    else:
        say(f"  ⚠ {compare_main}/pairwise_delta.csv 里没有 {v3_run} -> {main_run}")

    # ------------------------------------------------------------------ [7] Table S2 位置机制
    TS2 = pd.DataFrame()
    if pos_compare and pos_compare.get("enabled"):
        say(f"\n[7] Table S2：位置机制(以 {v3_run} 为底；{pos_compare['outdir']})")
        pdir = pos_compare["outdir"]
        if not os.path.exists(os.path.join(pdir, "delta_vs_reference.csv")):
            miss = [r for r in pos_compare["runs"] if not os.path.exists(os.path.join(results_root, r, "predictions_test.parquet"))]
            if miss:
                say(f"  缺导出 {miss}，跳过(这些实验的 checkpoint 在的话先 ./run_all.sh --export)")
            else:
                say("  第一次运行：调用 19 号 run_compare 重算(只算 perhead、不堆叠、不分层，CPU 约 5~10 分钟)……")
                t7 = time.time()
                try:
                    name = "_tflayout_19_compare_runs"
                    if name in sys.modules:
                        m19 = sys.modules[name]
                    else:
                        spec = importlib.util.spec_from_file_location(name, os.path.join(here, "19_compare_runs.py"))
                        m19 = importlib.util.module_from_spec(spec)
                        sys.modules[name] = m19
                        spec.loader.exec_module(m19)
                    m19.run_compare(results_root=results_root, runs=list(pos_compare["runs"]), references=(v3_run,),
                                    variants=("perhead",), seeds="common", offset_grid=tuple(offset_grid),
                                    n_boot=int(pos_compare.get("n_boot", 500)), stack=False, outdir=pdir,
                                    dense_target=m19.CONFIG.get("dense_target"), print_variants=("perhead",),
                                    skip_runs=(), headline_runs=(), verdict_variants=("perhead",),
                                    match_variant=True, pairwise_extra=True, strata=False,
                                    extra_pairs=tuple(pos_compare.get("extra_pairs", ())))
                    say(f"  19 号重算完成(用时 {time.time() - t7:.0f} 秒)")
                except Exception as e:  # noqa: BLE001 —— 这张是补充表，失败不影响其余
                    say(f"  ⚠ 19 号重算失败：{type(e).__name__}: {e}(补充表跳过，其余照常)")
        rowsS2 = []
        for run, lab_a, lab_t in pos_compare.get("labels", ()):
            q, k, src = find_delta(run, v3_run, "perhead", (pdir,))
            if q is not None:
                rowsS2 += [dict(r_, pair=f"{lab_a} vs v3") for r_ in delta_rows(q, k, src, lab_a, run)]
        for base_, run in pos_compare.get("extra_pairs", ()):
            q, k, src = find_delta(run, base_, "perhead", (pdir,))
            if q is not None:
                lb = dict((r, a) for r, a, _ in pos_compare.get("labels", ()))
                rowsS2 += [dict(r_, pair=f"{lb.get(run, run)} vs {lb.get(base_, base_)}")
                           for r_ in delta_rows(q, k, src, f"{lb.get(run, run)} vs {lb.get(base_, base_)}", run)]
        TS2 = pd.DataFrame(rowsS2)
        TS2.to_csv(os.path.join(outdir, "tableS2_position.csv"), index=False)
        if not TS2.empty:
            for pr in dict.fromkeys(TS2["pair"]):
                q = TS2[TS2["pair"] == pr]
                say(f"  {ascii_of(pr):42s}(k={int(q['k'].iloc[0])}): " + "  ".join(
                    f"{SHORT[r_.metric]} {f3(r_.delta, True)}{'†' if r_.test_sig else ''}{'‡' if r_.tot_sig else ''}"
                    for r_ in q.itertuples() if r_.metric in ("A_r_gene", "B_r", "B_r_within_tf", "C_auroc_down",
                                                               "C_auprc_down", "C_auprc_up")))
            say("  (这张表的模型是 v3 配置，方法里要写明；噪声 rms 用 v8 的零分布外推，只是量级参照)")

    # ------------------------------------------------------------------ [8] 自检
    say("\n[8] 自检(不同脚本算的同一个量必须一致；⚠ 的那项先查来源再写论文)")
    if mdl26 and hl is not None and t2b is not None:
        diffs = []
        for k26, kl in B26.items():
            a_ = t2b[(t2b["model"] == mdl26) & (t2b["metric"] == k26)]["value"]
            b_ = hl[(hl["run"] == main_run) & (hl["variant"] == "perhead") & (hl["metric"] == kl)]["ensemble"]
            if len(a_) and len(b_):
                diffs.append((k26, float(a_.iloc[0]) - float(b_.iloc[0])))
        mx = max((abs(d) for _, d in diffs), default=np.nan)
        say(f"  (a) 26 号模型行 vs 19 号第9批 [4] v8 perhead：最大差 {mx:.1e} " + ("✓" if mx < 2e-3 else "⚠ " + str(diffs)))
    if mdlA26 and mA is not None and t1b is not None:
        a_ = t1b[(t1b["model"] == mdlA26) & (t1b["gene_set"] == "test 全部")]["r"]
        b_ = mA[(mA["run"] == main_run) & (mA["split"] == "test") & (mA["stratum"] == "全部")]["r"]
        if len(a_) and len(b_):
            d_ = abs(float(a_.iloc[0]) - float(b_.iloc[0]))
            say(f"  (b) 26 号 Head A 模型 test 全部 vs 25 号第9批：差 {d_:.1e} " + ("✓" if d_ < 1e-6 else "⚠"))
    if not T4.empty:
        for run, lab_a, _ in ablations:
            a_ = T4[(T4["run"] == run) & (T4["metric"] == "A_r_gene") & (T4["scope"] == "test")]
            b_ = T4[(T4["run"] == run) & (T4["scope"] == "A全基因:test:网格")]
            if len(a_) and len(b_) and int(a_["k"].iloc[0]) == int(b_["k"].iloc[0]):
                d_ = abs(float(a_["delta"].iloc[0]) - float(b_["delta"].iloc[0]))
                say(f"  (c) {lab_a}：19 号 A_r(820 网格) Δ vs 25 号 test 网格 Δ 差 {d_:.1e} "
                    + ("✓" if d_ < 5e-3 else "⚠(超过 bf16/batch 组成造成的 ~1e-3 量级)"))
    if pgA is not None and f"{main_run}_ens" in pgA.columns:
        colsA = [c for c in pgA.columns if re.match(rf"^{re.escape(main_run)}_seed\d+$", c)]
        d_ = float(np.nanmax(np.abs(pgA[colsA].mean(axis=1).to_numpy() - pgA[f"{main_run}_ens"].to_numpy())))
        say(f"  (d) 25 号逐基因 csv：{main_run} 逐 seed 均值 vs _ens 列 最大差 {d_:.1e} " + ("✓" if d_ < 1e-6 else "⚠"))
    if t2b is not None and "B4_pair+TF先验" in set(t2b["model"]) and "B2_TF×结合先验" in set(t2b["model"]):
        g_ = {m_: t2b[(t2b["model"] == m_) & (t2b["metric"] == "C_AUCdn")]["value"].iloc[0]
              for m_ in ("B4_pair+TF先验", "B2_TF×结合先验")}
        say(f"  (e) 26 号 B4 ≥ B2(AUROC dn)：{g_['B4_pair+TF先验'] - g_['B2_TF×结合先验']:+.3f} "
            + ("✓" if g_["B4_pair+TF先验"] - g_["B2_TF×结合先验"] > -0.01 else "⚠"))

    # ------------------------------------------------------------------ [9] 规则核对 + paper_numbers.md
    say("\n[9] 规则核对(预先写下的规则，用这里的数重判；\"明显\"=|Δ|≥0.02 且 CI 不含0)")

    def dmb_get(base, met):
        if dmb is None:
            return np.nan, np.nan, np.nan
        q = dmb[(dmb["baseline"] == base) & (dmb["metric"] == met)]
        return (float(q["delta"].iloc[0]), float(q["ci_lo"].iloc[0]), float(q["ci_hi"].iloc[0])) if len(q) \
            else (np.nan, np.nan, np.nan)

    def clear(d, lo, hi, thr=0.02):
        return bool(np.isfinite(d) and abs(d) >= thr and (lo > 0 or hi < 0))

    verdicts = []
    if dmb is not None:
        a1, a2 = dmb_get("B4_pair+TF先验", "B_r|TF"), dmb_get("B4_pair+TF先验", "C_APdn")
        verdicts.append(("(a) 模型对 B4 在 B_r|TF、C_APdn 明显更好", clear(*a1) and clear(*a2),
                         f"B_r|TF {a1[0]:+.3f}[{a1[1]:+.3f},{a1[2]:+.3f}]  C_APdn {a2[0]:+.3f}[{a2[1]:+.3f},{a2[2]:+.3f}]"))
        worse = [(b_, m_) for b_ in ("B5_B4+实测WT表达", "B6_gbdt(同B5特征)", "B7a_gbdt+flat(无WT)", "B7b_gbdt+flat(+WT)", "B7a_bag(无WT)", "B7b_bag(+WT)") for m_ in B26
                 if clear(*dmb_get(b_, m_)) and dmb_get(b_, m_)[0] < 0]
        verdicts.append(("(c) B5/B6/B7 有指标明显强于模型(触发=要在讨论里写)", bool(worse), str(worse) if worse else "没有"))
    if t3b is not None and mdl26:
        best_b = {}
        for d_ in ("dn", "up"):
            q = t3b[(t3b["direction"] == d_) & t3b["model"].astype(str).str.startswith(("B3", "B4", "B5", "B6", "B7"))]
            if len(q):
                best_b[d_] = q.sort_values("pooled_auc").iloc[-1]["model"]
        ok_e, txt = True, []
        for d_, bm in best_b.items():
            q = t3b[t3b["model"] == f"Δ {mdl26} − {bm}"]
            q = q[q["direction"] == d_]
            if len(q):
                d0, lo0, hi0 = float(q["pooled_auc"].iloc[0]), float(q["pooled_ci_lo"].iloc[0]), float(q["pooled_ci_hi"].iloc[0])
                ok_e &= bool(d0 >= 0.03 and lo0 > 0)
                txt.append(f"{d_}: 对 {bm} {d0:+.3f}[{lo0:+.3f},{hi0:+.3f}]")
            else:
                ok_e = False
        verdicts.append(("(e) TF 内 AUROC 比 B3~B7(含 bag)最强的高 ≥0.03 且 CI 不含0", ok_e, "  ".join(txt)))
    if not T6.empty:   # 2026-10-02c 第14批：公平集成对照规则(跟 26 号文件头第14批同一套)
        def g6(key, met):
            q_ = T6[(T6["comparison"] == key) & (T6["metric"] == met)]
            return (float(q_["delta"].iloc[0]), float(q_["ci_lo"].iloc[0]), float(q_["ci_hi"].iloc[0])) if len(q_) \
                else (np.nan, np.nan, np.nan)

        for tag_, key_, lab_ in (("(i)", "(i) 集成−B7a_bag(对等集成)", "对等集成：模型集成−B7a_bag"),
                                 ("(j)", "(j) 单seed均值−B7a(单对单)", "单对单：模型单seed均值−B7a")):
            got_ = {m_: g6(key_, m_) for m_ in ("B_r|TF", "C_APdn", "wTF_dn")}
            if all(not np.isfinite(v_[0]) for v_ in got_.values()):
                continue
            ok3 = [m_ for m_, v_ in got_.items() if np.isfinite(v_[0]) and v_[0] >= 0.02 and v_[1] > 0]
            bad3 = [m_ for m_, v_ in got_.items() if np.isfinite(v_[0]) and v_[0] <= -0.02 and v_[2] < 0]
            verdicts.append((f"{tag_}+ {lab_} 在 B_r|TF/C_APdn/wTF_dn 里 ≥2 项明显更好【+ 口径：只含 test 抽样，成立也不据此写有增益，以 {tag_[:-1]}‡) 和 (m) 为准】", len(ok3) >= 2,
                             "  ".join(f"{m_} {v_[0]:+.3f}[{v_[1]:+.3f},{v_[2]:+.3f}]" for m_, v_ in got_.items())
                             + f"  明显更好{ok3} 明显更差{bad3}"))
        kk_ = g6("(k) 集成增益 B7a_bag−B7a", "C_APdn")
        if np.isfinite(kk_[0]):
            verdicts.append(("(k)+ B7a_bag−B7a 在 C_APdn ≥0.02 且 CI>0(集成对 GBDT 也有明显增益，对等性成立)【+ 口径：只含 test 抽样】", bool(kk_[0] >= 0.02 and kk_[1] > 0),
                             f"C_APdn {kk_[0]:+.3f}[{kk_[1]:+.3f},{kk_[2]:+.3f}]"))
        if "tot_lo" in T6.columns:   # 2026-10-02d 第15批：含训练随机性的 (i‡)(j‡) + 摘要版本 (m)
            def g6t(key, met):
                q_ = T6[(T6["comparison"] == key) & (T6["metric"] == met)]
                if not len(q_):
                    return (np.nan,) * 5
                r0 = q_.iloc[0]
                return tuple(float(r0[c_]) for c_ in ("delta", "tot_lo", "tot_hi", "cons_lo", "cons_hi"))

            n_it = None
            for tag_, key_, lab_ in (("(i‡)", "(i) 集成−B7a_bag(对等集成)", "对等集成：模型集成−B7a_bag"),
                                     ("(j‡)", "(j) 单seed均值−B7a(单对单)", "单对单：模型单seed均值−B7a")):
                got_ = {m_: g6t(key_, m_) for m_ in ("B_r|TF", "C_APdn", "wTF_dn")}
                if all(not np.isfinite(v_[0]) for v_ in got_.values()):
                    continue
                okt = [m_ for m_, v_ in got_.items() if np.isfinite(v_[1]) and v_[0] >= 0.02 and v_[1] > 0]
                okc = [m_ for m_ in okt if np.isfinite(got_[m_][3]) and got_[m_][3] > 0]
                if tag_ == "(i‡)":
                    n_it = (okt, okc)
                verdicts.append((f"{tag_} {lab_} 在 B_r|TF/C_APdn/wTF_dn 里 ≥2 项明显‡(合成区间含训练随机性)", len(okt) >= 2,
                                 "  ".join(f"{m_} {v_[0]:+.3f} 合成({v_[1]:+.3f},{v_[2]:+.3f}) 保守({v_[3]:+.3f},{v_[4]:+.3f})"
                                           for m_, v_ in got_.items()) + f"  明显‡{okt}(其中§也成立{okc})"))
            if n_it is not None:
                okt, okc = n_it
                var_ = "A" if len(okt) >= 2 else ("A-lite" if len(okt) == 1 else "B")
                verdicts.append(("(m) 摘要可以写\"对等集成下仍有增益\"(Variant A 或 A-lite；按 (i‡)，status 第9节(l))", var_ != "B",
                                 f"Variant {var_}：明显‡的项 {okt}" + ("" if var_ == "B" else f"(§ 也成立的 {okc}；摘要只点这些项的名字和数字)")
                                 + ("" if var_ != "B" else "——摘要和 Introduction 不写任何\"优于平铺特征\"")))
    if tb7 is not None and len(tb7):   # 2026-10-02d 第15批：B7 选中的轮数顶不顶边界
        sel_ = tb7[(tb7["head"].astype(str) == "B") | tb7["head"].astype(str).str.contains("选中")]
        edges_ = [f"{r_.model} {r_.head}={int(r_.selected)}" for r_ in sel_.itertuples() if tb(bool(str(r_.edge).strip().lower() in ("true", "1")))]
        verdicts.append(("(n) B7 选中的 Head B / Head C 轮数都不顶网格边界(基线调参收敛)", not edges_,
                         "  ".join(f"{r_.model} {r_.head}={int(r_.selected)}(val {float(r_.val_score):.4f})" for r_ in sel_.itertuples())
                         + (f"  顶边界：{edges_}" if edges_ else "")))
    if t5b is not None and len(t5b):
        ql_ = t5b[(t5b["stack"].astype(str) == "STACK(模型+B7b_bag)") & (t5b["reference"].astype(str) == "B7b_bag(+WT)")]
        if len(ql_):
            ql_ = ql_.set_index("metric")
            got_l = {m_: (float(ql_.loc[m_, "delta"]), float(ql_.loc[m_, "ci_lo"]), float(ql_.loc[m_, "ci_hi"])) for m_ in ("B_r|TF", "C_APdn") if m_ in ql_.index}
            verdicts.append(("(l) 堆叠(模型+B7b_bag)−B7b_bag 在 B_r|TF 或 C_APdn 明显>0(模型带来最强基线之外的信息，可写\"互补\")",
                             any(v_[0] >= 0.02 and v_[1] > 0 for v_ in got_l.values()),
                             "  ".join(f"{m_} {v_[0]:+.3f}[{v_[1]:+.3f},{v_[2]:+.3f}]" for m_, v_ in got_l.items())))
    if not T3.empty:
        strong = [m_ for m_ in ("A3_layout+token_分块ridge", "A3k_layout+6mer_分块ridge", "A4_gbdt(layout+token)",
                                "A4k_gbdt(layout+6mer)") if m_ in set(T3["model"])]
        if strong:
            bm = max(strong, key=lambda m_: float(T3[T3["model"] == m_]["test 全部"].iloc[0]))
            r_ = T3[T3["model"] == bm].iloc[0]
            verdicts.append(("(d) Head A 对最强浅层基线 Δ≥0.02(全部基因)", bool(r_["test 全部_d"] >= 0.02),
                             f"对 {bm}: 全部 {r_['test 全部_d']:+.3f}[{r_['test 全部_dlo']:+.3f},{r_['test 全部_dhi']:+.3f}]  "
                             f"网格外 {r_.get('test 网格外_d', np.nan):+.3f}[{r_.get('test 网格外_dlo', np.nan):+.3f},"
                             f"{r_.get('test 网格外_dhi', np.nan):+.3f}]"))
    verdicts += rules_stack(t5b, clear)   # 2026-10-02e 第16批：每个堆叠 x 每个参照一行，带标签
    if t4b is not None:
        q = t4b[t4b["baseline"] == "B4_pair+TF先验"].set_index(["stratum", "metric"])["delta"]
        try:
            gin = [q.loc[("D∈L_g", m_)] for m_ in ("B_r|TF", "C_APdn")]
            gout = [q.loc[("D∉L_g", m_)] for m_ in ("B_r|TF", "C_APdn")]
            verdicts.append(("(g) 模型对 B4 的增量主要在 D∈L_g 层", bool(np.mean(gin) > np.mean(gout)),
                             f"D∈L_g B_r|TF {gin[0]:+.3f} C_APdn {gin[1]:+.3f}；D∉L_g {gout[0]:+.3f} / {gout[1]:+.3f}"))
        except KeyError:
            pass
    if not T4.empty:
        def t4(run, met, scope="test"):
            q = T4[(T4["run"] == run) & (T4["metric"] == met) & (T4["scope"] == scope)]
            return q.iloc[0] if len(q) else None
        r_lay = [t4("v10_abl_no_layout", m_) for m_ in ("A_r_gene", "B_r", "B_r_within_tf", "C_auprc_down")]
        if all(x is not None for x in r_lay):
            half = (-.358 / 2, -.313 / 2, -.341 / 2, -.196 / 2)
            verdicts.append(("[消融] no_layout 降幅 ≥ 第4批的一半，且合成区间不含0",
                             all(x["delta"] <= h and x["tot_sig"] for x, h in zip(r_lay, half)),
                             "  ".join(f"{SHORT[x['metric']]} {x['delta']:+.3f}{'‡' if x['tot_sig'] else ''}" for x in r_lay)))
        r_ko = [t4("v10_abl_no_knockout", m_) for m_ in ("C_auprc_down", "C_auprc_up", "B_r_within_tf", "A_r_gene")]
        if all(x is not None for x in r_ko):
            verdicts.append(("[消融] no_knockout：C_APdn 变差超出合成区间(\"删位点帮 Head C\")",
                             bool(r_ko[0]["delta"] < 0 and r_ko[0]["tot_sig"]),
                             "  ".join(f"{SHORT[x['metric']]} {x['delta']:+.3f}{'†' if x['test_sig'] else ''}"
                                       f"{'‡' if x['tot_sig'] else ''}(k={int(x['k'])})" for x in r_ko)
                             + "  (A_r 是阴性对照：设计上不该动)"))
        if r_ko[0] is not None and "cons_sig" in T4.columns:  # 2026-10-02a 第12批
            x0 = r_ko[0]
            verdicts.append(("[消融·第12批] no_knockout 的 C_APdn 在保守合成区间(§，不外推)下也不含0",
                             bool(x0["delta"] < 0 and tb(x0.get("cons_sig", False))),
                             f"C_APdn {x0['delta']:+.3f} 合成({f3(x0['tot_lo'], True)},{f3(x0['tot_hi'], True)}) 保守("
                             f"{f3(x0.get('cons_lo', np.nan), True)},{f3(x0.get('cons_hi', np.nan), True)})  逐seed "
                             f"{int(x0['seeds_better']) if np.isfinite(x0['seeds_better']) else '?'}/{int(x0['k'])} 变好  "
                             f"符号检验双侧p={f3(x0.get('sign_p', np.nan))}"))
            robust, fragile = [], []
            for x in T4[(T4["run"] == "v10_abl_no_knockout") & (T4["scope"] == "test")].itertuples():
                if not tb(x.tot_sig) or x.metric == "A_r_gene":
                    continue
                nb = x.seeds_better
                same = (x.k - nb) if x.delta < 0 else nb
                ok_x = tb(x.cons_sig) and np.isfinite(nb) and same >= x.k - 1
                (robust if ok_x else fragile).append(
                    f"{SHORT[x.metric]} {x.delta:+.3f}(同向 {int(same) if np.isfinite(nb) else '?'}/{int(x.k)}"
                    f"{'' if tb(x.cons_sig) else '，§不成立'})")
            verdicts.append(("[消融·第12批] no_knockout 过‡的指标里，同时过§且逐seed≥k−1同向的(可写成结论)", bool(robust),
                             "稳：" + ("、".join(robust) or "无") + "；不稳(只写\"在 test 抽样意义下\")：" + ("、".join(fragile) or "无")))
        if not T5.empty:
            q = T5[(T5["against"] == "w/o site deletion") & (T5["metric"] == "C_auprc_down")].set_index("stratum")
            if {"D∈L_g", "D∉L_g"} <= set(q.index):
                verdicts.append(("[消融] no_knockout 的 C_APdn 变差只在 D∈L_g(D∉L_g 合成区间含0)",
                                 bool(q.loc["D∈L_g", "tot_sig"] and not q.loc["D∉L_g", "tot_sig"]),
                                 f"D∈L_g {q.loc['D∈L_g', 'delta']:+.3f}{'‡' if q.loc['D∈L_g', 'tot_sig'] else ''}  "
                                 f"D∉L_g {q.loc['D∉L_g', 'delta']:+.3f}{'‡' if q.loc['D∉L_g', 'tot_sig'] else ''}"))
        r_cis = [t4("v10_abl_no_cis", m_) for m_ in ("A_r_gene", "C_auprc_down", "C_macro_f1_offset")]
        if all(x is not None for x in r_cis):
            verdicts.append(("[消融] no_cis 有任何一项超出合成区间", any(x["tot_sig"] for x in r_cis),
                             "  ".join(f"{SHORT[x['metric']]} {x['delta']:+.3f}{'‡' if x['tot_sig'] else ''}" for x in r_cis)))
        r_la = [t4("v10_lambda_a03", m_) for m_ in ("B_r_within_tf", "C_auprc_down", "A_r_gene")]
        if all(x is not None for x in r_la):
            verdicts.append(("[消融] λ_A=0.3：B_r|TF 或 C_APdn 明显更好(2/2)且 A_r 降幅 ≤0.03",
                             bool(any(x["delta"] >= 0.02 and x["tot_sig"] and x["seeds_better"] == x["k"]
                                      for x in r_la[:2]) and r_la[2]["delta"] >= -0.03),
                             "  ".join(f"{SHORT[x['metric']]} {x['delta']:+.3f}{'‡' if x['tot_sig'] else ''}" for x in r_la)))
    pr_csv = os.path.join(probe_dir, "rule_checks_probe.csv") if probe_dir else None  # 2026-10-02a 第12批：28 号
    qpr = rd(pr_csv)
    if qpr is not None and len(qpr):
        for r_ in qpr.itertuples():
            verdicts.append((f"[28号] {r_.rule}", str(r_.holds).strip().lower() in ("true", "1"), str(r_.numbers)))
    elif probe_dir:
        say(f"  (没有 {pr_csv}：28 号还没跑；./run_all.sh --probe --tables --skip-arch-selftest 后这里会多出 [28号] 几行)")
    sp_csv = os.path.join(spec_dir, "rule_checks_spec.csv") if spec_dir else None  # 2026-10-02b 第13批：29 号
    qsp = rd(sp_csv)
    if qsp is not None and len(qsp):
        for r_ in qsp.itertuples():
            verdicts.append((f"[29号] {r_.rule}", str(r_.holds).strip().lower() in ("true", "1"), str(r_.numbers)))
    elif spec_dir:
        say(f"  (没有 {sp_csv}：29 号还没跑；./run_all.sh --spec-control --tables 后这里会多出 [29号] 几行)")
    for nm, ok_, txt in verdicts:
        say(f"  {'成立' if ok_ else '不成立'}  {nm}：{txt}")
    pd.DataFrame([dict(rule=nm, holds=ok_, numbers=txt) for nm, ok_, txt in verdicts]).to_csv(
        os.path.join(outdir, "rule_checks.csv"), index=False)

    # ------------------------------------------------------------------ [9b] 短文(4~5 页)专用：紧凑表 + 自动摘要(第16批 v6)
    say("\n[9b] 短文(ICBCB/ICBBT 4~5 页)专用：tables_short.tex/.docx/.txt(两张紧凑表) + abstract_draft.md(按 (m) 选版本、自动填数)")
    short_md = []
    try:
        fixed_ = dict(fixed_facts or dict(n_tf=178, n_sites=214329, n_promoters=5373, source="默认常量"))
        n_seeds_ = 5
        if not T1.empty and "n_seeds" in T1.columns and T1["n_seeds"].notna().any():
            n_seeds_ = int(pd.to_numeric(T1["n_seeds"], errors="coerce").dropna().iloc[0])
        spec_df_ = rd(os.path.join(spec_dir, "contrasts_spec.csv")) if spec_dir else None
        short_tex, short_tables, short_plain = [], [], []
        for tbl_ in (short_main_table(T2, mdl26, T6, t6v, facts["test"], n_seeds_),
                     short_mech_table(T4, T5, ablations, spec_df_, facts["test"].get("d_in_lg"))):
            if tbl_ is None:
                continue
            short_plain.append(tbl_["caption_txt"])
            short_plain += ["    " + x_ for x_ in render_short_table(tbl_, tex_table, short_tex)]
            short_plain.append("")
            short_tables.append(tbl_)
        for x_ in short_plain:
            say("  " + x_)
        if short_tex:
            with open(os.path.join(outdir, "tables_short.tex"), "w", encoding="utf-8") as fh:
                fh.write("% 27_paper_tables.py v6 自动生成(4~5 页短文用，单栏 table)；需要 \\usepackage{booktabs,graphicx,amssymb}。\n\n" + "\n".join(short_tex))
            with open(os.path.join(outdir, "tables_short.txt"), "w", encoding="utf-8") as fh:
                fh.write("\n".join(short_plain) + "\n")
            ok_docx = write_short_docx(os.path.join(outdir, "tables_short.docx"), short_tables)
            say("  写出 tables_short.tex、tables_short.txt" + ("、tables_short.docx(Word 版，粘进会议模板)" if ok_docx
                                                              else "(没有 python-docx，跳过 .docx：pip install python-docx 后重跑 --tables)"))
        else:
            say("  (没有足够的数据生成短文表：先 --baselines 再 --tables)")
        ab_ = build_abstract(T2, mdl26, T4, T5, T6, t5b, facts, spec_df_, fixed_, n_seeds_)
        say(f"  摘要版本 (m) = Variant {ab_['variant']}；{ab_['words']} 词；flags {len(ab_['flags'])} 条")
        for f_ in ab_["flags"]:
            say(f"    ⚠ {f_}")
        with open(os.path.join(outdir, "abstract_draft.md"), "w", encoding="utf-8") as fh:
            fh.write(f"# 摘要草稿(27 号 v6 自动生成，{time.strftime('%Y-%m-%d %H:%M')})\n\n")
            fh.write(f"摘要版本(m)：**Variant {ab_['variant']}**；{ab_['words']} 词。数字只从本目录的 csv 来，常量标明来源。\n\n")
            fh.write(ab_["text"] + "\n\n## 条件检查(flags 为空 = 每句话的前提都满足)\n\n")
            fh.write("\n".join(f"- ⚠ {f_}" for f_ in ab_["flags"]) if ab_["flags"] else "- 全部满足")
            fh.write("\n\n## 数字来源\n\n" + "\n".join(f"- {a_} ← {b_}" for a_, b_ in ab_["sources"]) + "\n")
            fh.write("\n## 没有自动核对的一句\n\n- We also document failure modes: unseen TFs and spacing below ~100 bp. "
                     "(定性，来自 status 7.12/7.14 的 LTO 和 7.13/7.16 的位置分辨率；改了这两部分结论要手动改)\n")
        short_md.append("\n## 摘要与短文表 — abstract_draft.md / tables_short.tex\n")
        short_md.append(f"- 摘要版本 (m)：Variant {ab_['variant']}；常量 {fixed_['n_tf']} 个 TF、{fixed_['n_sites']:,} 个位点、{fixed_['n_promoters']:,} 个启动子"
                        f"(来源：{fixed_.get('source', '')}；不是本脚本算的)")
        for a_, b_ in ab_["sources"]:
            short_md.append(f"- {a_} ← {b_}")
    except Exception as e_:  # noqa: BLE001 —— [9b] 失败不影响前面的输出
        import traceback
        say(f"  ⚠ [9b] 失败(不影响上面的表、规则、paper_numbers.md)：{type(e_).__name__}: {e_}")
        say(traceback.format_exc())

    # paper_numbers.md：写作时只从这里抄
    md.append(f"# 论文数字清单(27 号自动生成，{time.strftime('%Y-%m-%d %H:%M')})\n")
    md.append("每个数字后面的方括号是 95% CI(按基因整群 bootstrap，只含 test 抽样)；来源路径见各节。不要从别处手抄。\n")
    md.append("## 数据规模(从主模型 17 号导出直接数)\n")
    for sp_ in ("val", "test"):
        f_ = facts[sp_]
        md.append(f"- {sp_}: {f_['rows']} 对 = {f_['genes']} 基因 × {f_['tfs']} 个被耗竭 TF；显著 {f_['sig']}(down {f_['down']} / "
                  f"up {f_['up']})；D∈L_g {f_['d_in_lg']:.1%}")
    md.append(f"\n## Table 1 主结果({main_run} 5-seed 集成，perhead)— {compare_main}/headline_all_seeds.csv 等\n")
    for r_ in T1.itertuples() if not T1.empty else []:
        if r_.variant in ("perhead", "perhead+stack"):
            md.append(f"- {r_.block} / {r_.variant} / {r_.metric}: **{f3(r_.value)}** [{f3(r_.ci_lo)}, {f3(r_.ci_hi)}]  "
                      f"(`{r_.source}`)")
    md.append(f"\n## Table 2/3 对基线 — {baselines}/\n")
    if dmb is not None:
        for b_ in ("B4_pair+TF先验", "B6_gbdt(同B5特征)", "B7a_gbdt+flat(无WT)", "B7b_gbdt+flat(+WT)"):
            q = dmb[dmb["baseline"] == b_]
            md.append(f"- 模型 − {b_}: " + "；".join(f"{r_.metric} {r_.delta:+.3f} [{r_.ci_lo:+.3f}, {r_.ci_hi:+.3f}]"
                                                     for r_ in q.itertuples()))
    if t3b is not None:
        q = t3b[t3b["model"].astype(str).str.startswith("Δ")]
        for r_ in q.itertuples():
            md.append(f"- TF 内 AUROC(合并) {r_.model} {r_.direction}: {r_.pooled_auc:+.3f} [{r_.pooled_ci_lo:+.3f}, "
                      f"{r_.pooled_ci_hi:+.3f}]")
    md.append(f"\n## Table 4 消融(在 {main_run} 上)— 见 table4_ablation_long.csv 的 source 列\n")
    md.append("† = bootstrap CI 不含 0；‡ = 合成区间(再加训练随机性)不含 0；§ = 保守合成区间(k≥3 时不外推)不含 0。论文里只把 ‡ 的写成"
              "\"显著\"(‡ 和 § 都成立的最稳)，只有 † 的写\"在 test 抽样意义下\"。符号检验 p 是逐 seed 同向的双侧精确二项检验。\n")
    for r_ in T4.itertuples() if not T4.empty else []:
        if (r_.scope == "test" or r_.scope.startswith("A全基因:test")) and np.isfinite(r_.delta):
            md.append(f"- {r_.label} (k={r_.k}) {r_.scope} {SHORT.get(r_.metric, r_.metric)}: Δ {r_.delta:+.3f} "
                      f"[{r_.ci_lo:+.3f}, {r_.ci_hi:+.3f}]{'†' if r_.test_sig else ''}  合成 [{f3(r_.tot_lo, True)}, "
                      f"{f3(r_.tot_hi, True)}]{'‡' if r_.tot_sig else ''}"
                      + (f"  保守 [{f3(r_.cons_lo, True)}, {f3(r_.cons_hi, True)}]{'§' if tb(r_.cons_sig) else ''}"
                         if int(r_.k) > 2 else "")
                      + (f"  逐seed {int(r_.seeds_better)}/{int(r_.k)} 变好、符号检验 p={r_.sign_p:.3f}"
                         if np.isfinite(r_.sign_p) else ""))
    md.append("\n## Table 5 分层 — table5_strata_long.csv\n")
    for r_ in T5.itertuples() if not T5.empty else []:
        md.append(f"- {r_.kind} {r_.against} {r_.stratum} {SHORT.get(r_.metric, r_.metric)}: Δ {r_.delta:+.3f} "
                  f"[{r_.ci_lo:+.3f}, {r_.ci_hi:+.3f}]{'†' if r_.test_sig else ''}"
                  + ("‡" if tb(getattr(r_, "tot_sig", False)) else ""))
    md.append("\n## Table 6 公平集成对照(26 号 [2d]；‡/§ 要 26 号 v6)— table6_fair_ensemble.csv / table6_values.csv\n")
    if not T6.empty:
        has_tot6 = "tot_sig" in T6.columns
        md.append("† = bootstrap CI 不含 0(只含 test 抽样)" + ("；‡ = 合成区间(再加模型换 seed、GBDT 换成员的重训随机性)不含 0；§ = 保守区间(模型集成不外推、"
                                                            "B7 单模型也计重训噪声)不含 0。论文里\"对等集成下的增益\"只写 ‡ 成立的项" if has_tot6 else "")
                  + "；单 seed 均值的 seed 间标准差在 table6_values.csv。\n")
        for r_ in T6.itertuples():
            md.append(f"- {r_.comparison} / {r_.metric}: Δ {r_.delta:+.3f} [{r_.ci_lo:+.3f}, {r_.ci_hi:+.3f}]"
                      f"{'†' if (r_.ci_lo > 0 or r_.ci_hi < 0) else ''}"
                      + (f"  合成 [{f3(getattr(r_, 'tot_lo', np.nan), True)}, {f3(getattr(r_, 'tot_hi', np.nan), True)}]"
                         f"{'‡' if tb(getattr(r_, 'tot_sig', False)) else ''}  保守 [{f3(getattr(r_, 'cons_lo', np.nan), True)}, "
                         f"{f3(getattr(r_, 'cons_hi', np.nan), True)}]{'§' if tb(getattr(r_, 'cons_sig', False)) else ''}"
                         if has_tot6 else "")
                      + f"  ({r_.a_value:.3f} vs {r_.b_value:.3f})")
    else:
        md.append("(没有 26 号的 table6 输出)")
    md.append("\n## Table 6 的零分布(26 号 v6 table6_null.csv；两次独立重训之间差值的均方根)\n")
    if t6n is not None and len(t6n):
        for r_ in t6n.itertuples():
            if r_.metric in ("B_r|TF", "C_APdn", "C_APup", "C_AUCdn", "wTF_dn", "wTF_up"):
                md.append(f"- {r_.source} k={r_.k} {r_.metric}: rms {r_.rms:.4f}({r_.n_pairs} 组；{r_.members})")
    else:
        md.append("(没有 table6_null.csv：26 号还是 v5，或关了 fair_noise)")
    md.append("\n## Methods：B7 的超参与选中轮数(26 号 v6 table_b7_tuning.csv；HistGradientBoosting，early_stopping=False，轮数在 val 上选)\n")
    if tb7 is not None and len(tb7):
        for r_ in tb7.itertuples():
            md.append(f"- {r_.model} Head {r_.head}: 选中 {int(r_.selected)} 轮(val {'r' if str(r_.head) == 'B' else 'AUPRC 均值'} "
                      f"{float(r_.val_score):.4f}；{'⚠顶边界' if str(r_.edge).strip().lower() in ('true', '1') else '不顶边界'})；候选 {r_.grid}；"
                      f"learning_rate {r_.learning_rate}、max_leaf_nodes {int(r_.max_leaf_nodes)}、min_samples_leaf {int(r_.min_samples_leaf)}、"
                      f"特征 {int(r_.n_features)} 维")
    else:
        md.append("(没有 table_b7_tuning.csv：26 号还是 v5)")
    md.append("\n## 噪声底(训练随机性，v8 5 个 seed)— null_seed_variability.csv\n")
    for r_ in null_df.itertuples() if not null_df.empty else []:
        if r_.variant == "perhead" and r_.k == 2 and (r_.scope == "test" or r_.scope == "A全基因:test:全部"):
            md.append(f"- {r_.scope} {SHORT.get(r_.metric, r_.metric)}: 2 对 2 rms {r_.rms:.4f}(最大 {r_.max_abs:.4f}，{r_.n_pairs} 组)")
    md.extend(short_md)
    md.append("\n## 规则核对 — rule_checks.csv\n")
    for nm, ok_, txt in verdicts:
        md.append(f"- {'成立' if ok_ else '不成立'}：{nm} —— {txt}")
    for fn_, buf, dir_, who in (("paper_numbers_probe.md", md, probe_dir, "28"), ("tables_probe.tex", tex, probe_dir, "28"),
                                ("paper_numbers_spec.md", md, spec_dir, "29"), ("tables_spec.tex", tex, spec_dir, "29")):
        p_ = os.path.join(dir_, fn_) if dir_ else None   # 2026-10-02a 第12批并入 28 号；2026-10-02b 第13批并入 29 号
        if p_ and os.path.exists(p_):
            with open(p_, encoding="utf-8") as fh:
                buf.append(fh.read())
            say(f"  并入 {who} 号 {p_}")
    with open(os.path.join(outdir, "paper_numbers.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(md) + "\n")
    with open(os.path.join(outdir, "tables.tex"), "w", encoding="utf-8") as fh:
        fh.write("% 27_paper_tables.py 自动生成；需要 \\usepackage{booktabs,graphicx,amssymb}。数字来源见同目录 *.csv 的 source 列。\n\n"
                 + "\n".join(tex))

    # ------------------------------------------------------------------ [10] 图
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        if not T4.empty:
            mets = ["A_r_gene", "B_r", "B_r_within_tf", "C_auroc_down", "C_auprc_down", "C_auprc_up", "C_macro_f1_offset"]
            labs = [a for _, a, _ in ablations if a in set(T4["label"])]
            fig, ax = plt.subplots(1, len(labs), figsize=(4.2 * len(labs), 3.6), squeeze=False)
            for j, lab_a in enumerate(labs):
                q = T4[(T4["label"] == lab_a) & (T4["scope"] == "test")].set_index("metric").reindex(mets)
                y_ = np.arange(len(mets))[::-1]
                ax[0, j].errorbar(q["delta"], y_, xerr=[np.clip(q["delta"] - q["tot_lo"], 0, None),
                                                        np.clip(q["tot_hi"] - q["delta"], 0, None)],
                                  fmt="none", ecolor="0.6", elinewidth=1, capsize=2)
                ax[0, j].errorbar(q["delta"], y_, xerr=[np.clip(q["delta"] - q["ci_lo"], 0, None),
                                                        np.clip(q["ci_hi"] - q["delta"], 0, None)],
                                  fmt="o", color="k", ms=3, elinewidth=2.2, capsize=0)
                ax[0, j].axvline(0, color="r", lw=.6)
                ax[0, j].set_yticks(y_)
                ax[0, j].set_yticklabels([SHORT[m_] for m_ in mets], fontsize=8)
                ax[0, j].set_title(f"{lab_a} (k={int(q['k'].dropna().iloc[0]) if q['k'].notna().any() else '?'})",
                                   fontsize=9)
                ax[0, j].set_xlabel("Delta vs reference", fontsize=8)
            fig.suptitle("Ablations: thick = test bootstrap CI, thin = + seed variability", fontsize=9)
            fig.tight_layout()
            fig.savefig(os.path.join(outdir, "fig", "ablation_forest.png"), dpi=200)
            plt.close(fig)
        if not T2.empty and "wtf_dn" in T2.columns:
            q = T2[~T2["model"].astype(str).str.startswith(("ORACLE", "STACK"))]
            fig, ax = plt.subplots(figsize=(7.5, 3.4))
            x_ = np.arange(len(q))
            for j, d_ in enumerate(("dn", "up")):
                v = q[f"wtf_{d_}"].to_numpy(np.float64)
                ax.bar(x_ + 0.38 * j, v, 0.38, label=d_,
                       yerr=[np.clip(v - q[f"wtf_{d_}_lo"].to_numpy(np.float64), 0, None),
                             np.clip(q[f"wtf_{d_}_hi"].to_numpy(np.float64) - v, 0, None)], capsize=2)
            ax.axhline(0.5, color="k", lw=.6, ls="--")
            ax.set_xticks(x_ + 0.19)
            ax.set_xticklabels([ascii_of(s)[:18] for s in q["label"]], rotation=25, ha="right", fontsize=7)
            ax.set_ylim(0.4, max(0.85, float(np.nanmax(q[["wtf_dn", "wtf_up"]].to_numpy())) + 0.05))
            ax.set_ylabel("pooled within-TF AUROC")
            ax.legend(fontsize=7)
            fig.tight_layout()
            fig.savefig(os.path.join(outdir, "fig", "within_tf_auroc.png"), dpi=200)
            plt.close(fig)
        if not T5.empty:
            q = T5[T5["kind"] == "模型−基线"]
            if len(q):
                mets = ["B_r", "B_r_within_tf", "C_auprc_down", "C_auprc_up"]
                fig, ax = plt.subplots(1, 2, figsize=(9, 3.2), sharey=True)
                for j, st in enumerate(("D∈L_g", "D∉L_g")):
                    qq = q[q["stratum"] == st]
                    bases = list(dict.fromkeys(qq["against"]))
                    w = 0.8 / max(len(bases), 1)
                    for i, b_ in enumerate(bases):
                        z = qq[qq["against"] == b_].set_index("metric").reindex(mets)
                        ax[j].bar(np.arange(len(mets)) + i * w, z["delta"], w, label=ascii_of(NAME.get(b_, (b_,))[0]),
                                  yerr=[np.clip(z["delta"] - z["ci_lo"], 0, None), np.clip(z["ci_hi"] - z["delta"], 0, None)],
                                  capsize=2)
                    ax[j].axhline(0, color="k", lw=.6)
                    ax[j].set_xticks(np.arange(len(mets)) + 0.4 - w / 2)
                    ax[j].set_xticklabels([SHORT[m_] for m_ in mets], fontsize=8)
                    ax[j].set_title(("D in L_g" if st == "D∈L_g" else "D not in L_g") + ": ours - baseline", fontsize=9)
                ax[0].legend(fontsize=7)
                fig.tight_layout()
                fig.savefig(os.path.join(outdir, "fig", "strata_vs_baselines.png"), dpi=200)
                plt.close(fig)
    except Exception as e:  # noqa: BLE001 —— 画图失败不影响表格
        say(f"  (画图跳过: {type(e).__name__}: {e})")

    say(f"\n写出 {outdir}/：summary.txt、paper_numbers.md、tables.tex、table1_main.csv、table2_head_bc.csv、table3_head_a.csv、"
        f"table4_ablation_long.csv、table5_strata_long.csv、tableS1_v8_vs_v3.csv、tableS2_position.csv、"
        f"null_seed_variability.csv、rule_checks.csv、tables_short.tex/.docx/.txt、abstract_draft.md、fig/*.png(28/29 号的表/数字/规则已并入，如有)  (用时 {(time.time() - t_all) / 60:.1f} 分钟)")
    with open(os.path.join(outdir, "summary.txt"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return dict(null=null_df, table4=T4, table5=T5, verdicts=verdicts)


if __name__ == "__main__":
    run_paper_tables(**CONFIG)